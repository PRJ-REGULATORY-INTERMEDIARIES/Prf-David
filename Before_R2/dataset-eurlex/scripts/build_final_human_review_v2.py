"""Build the actor-in-role final human-review package (production v1.3).

This is a transformation of the extant, researcher-adjudicated/reconstructed
relationship matrix.  It deliberately does not discover, recode, delete, or
overwrite any historical production artifact.
"""
from __future__ import annotations

import csv
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
CASES = ROOT / "cases"
SOURCE = CASES / "FINAL_DATASET_FOR_HUMAN_REVIEW.csv"
UNCERTAINTIES = CASES / "R2_HUMAN_VALIDATION_uncertainties.csv"
OUT = CASES / "final_human_review_v2"
METHOD_VERSION = "production_v1.3"
PROMPT_VERSION = "production_v1.2 PRIMARY_INTERPRETIVE_READER (historical source run)"
DECISION_COLUMN = "PESQUISADOR: APPROVE / CHANGE / EXCLUDE / UNRESOLVED"

ACTS = {
    "case_01": {
        "title": "Regulation (EU) 2021/1119 — European Climate Law",
        "celex": "32021R1119",
    },
    "case_02": {
        "title": "Regulation (EU) 2023/956 — Carbon Border Adjustment Mechanism (CBAM)",
        "celex": "32023R0956",
    },
    "case_03": {
        "title": "Regulation (EU) 2023/1115 — Deforestation-free products Regulation (EUDR)",
        "celex": "32023R1115",
    },
}

# These are explicit, auditable classifications applied only to accepted I rows.
ORIENTATION = {
    "REL_C01_003": "regulator-facing",
    "REL_C01_007": "regulator-facing",
    "REL_C01_008": "regulator-facing",
    "REL_C02_001": "bidirectional",
    "REL_C02_002": "target-facing",
    "REL_C02_003": "bidirectional",
    "REL_C02_009": "regulator-facing",
    "REL_C02_010": "regulator-facing",
    "REL_C03_002": "bidirectional",
    "REL_C03_004": "regulator-facing",
    "REL_C03_006": "bidirectional",
    "REL_C03_007": "regulator-facing",
}


def read_csv(path: Path, delimiter: str = ",") -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle, delimiter=delimiter))


def write_csv(path: Path, rows: list[dict[str, str]], fields: list[str]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="raise")
        writer.writeheader()
        writer.writerows(rows)
    # Read-back validation protects against another delimiter/quoting failure.
    checked = read_csv(path)
    if len(checked) != len(rows) or (checked and list(checked[0]) != fields):
        raise RuntimeError(f"CSV validation failed: {path}")


def values(value: str) -> list[str]:
    return [item.strip() for item in value.split(";") if item.strip()]


def actor_name(value: str, role: str) -> str:
    """Remove the coding suffix while retaining an auditable verbatim source field."""
    marker = {
        "R": r"\s+\((?:considered\s+)?regulator\b",
        "I": r"\s+\((?:candidate\s+)?intermediary\b",
        "T": r"\s+\(target\b",
    }[role]
    return re.split(marker, value, maxsplit=1, flags=re.IGNORECASE)[0].strip()


def normalize_actor(verbatim: str) -> tuple[str, str]:
    """Return normalized name and a logged, intentionally conservative reason."""
    aliases = {
        "European Commission": "European Commission",
        "European Parliament and Council": "European Parliament and Council",
        "European Environment Agency (EEA)": "European Environment Agency (EEA)",
        "European Scientific Advisory Board on Climate Change (Advisory Board)": "European Scientific Advisory Board on Climate Change (Advisory Board)",
        "Energy Union Committee": "Energy Union Committee",
        "CBAM Committee": "CBAM Committee",
        "Comitology committee, Article 36": "Comitology committee, Article 36",
        "Member State competent authorities": "Member State competent authorities",
        "Member States": "Member States",
        "national accreditation bodies": "National accreditation bodies",
        "accredited verifiers": "Accredited verifiers",
        "customs authorities": "Customs authorities",
        "Experts designated by each Member State": "Experts designated by each Member State",
        "FLEGT licensing authorities/scheme (operating under Regulation (EC) No 2173/2005 Voluntary Partnership Agreements)": "FLEGT licensing authorities/scheme",
    }
    if verbatim in aliases:
        normalized = aliases[verbatim]
        reason = "Explicit canonicalization of a role-bearing label; no cross-actor merge."
    else:
        normalized = verbatim
        reason = "Retained verbatim: no safe alias merge was required."
    return normalized, reason


def actor_metadata(actor: str) -> tuple[str, str]:
    lower = actor.lower()
    if actor == "European Parliament and Council":
        return "EU legislature", "EU"
    if "commission" in lower:
        return "EU institution", "EU"
    if "agency" in lower or "advisory board" in lower:
        return "EU body", "EU"
    if "committee" in lower:
        return "EU comitology committee", "EU"
    if "member state competent" in lower:
        return "competent authority", "Member State"
    if actor == "Member States":
        return "governmental actor class", "Member State"
    if "customs" in lower:
        return "customs authority", "Member State / EU border"
    if "accreditation" in lower:
        return "accreditation body", "Member State"
    if "verifier" in lower:
        return "conformity-assessment actor", "NOT_CODED"
    if "expert" in lower:
        return "expert group", "Member State-designated"
    if "flegt" in lower:
        return "licensing authority/scheme", "NOT_CODED"
    if any(term in lower for term in ("declarant", "operator", "trader", "sector")):
        return "regulated economic actor/class", "NOT_CODED"
    if "natural or legal persons" in lower:
        return "person/organisation class", "NOT_CODED"
    return "NOT_CODED", "NOT_CODED"


def relation_status(row: dict[str, str]) -> str:
    if row["relation_type"] == "uncertain":
        return "UNCERTAIN"
    if row["relation_type"] == "not_supported":
        return "NOT_SUPPORTED_RESEARCHER_CONFIRMED"
    if row["provenance"].startswith("promoted_from_uncertainty_log"):
        return "RESEARCHER_MODIFIED"
    return "RESEARCHER_APPROVED"


def decision_metadata(row: dict[str, str], uncertainty_index: dict[str, dict[str, str]]) -> dict[str, str]:
    promoted = re.search(r"(case_\d\d-U\d\d)", row["provenance"])
    if promoted:
        source = uncertainty_index[promoted.group(1)]
        rationale = source["Correção final"] or "NOT_RECORDED"
        return {
            "human_decision": source["Decisão do pesquisador"] or "NOT_RECORDED",
            "human_decision_source": f"cases/R2_HUMAN_VALIDATION_uncertainties.csv:{promoted.group(1)}",
            "human_rationale": rationale,
            "decision_date": "NOT_RECORDED",
            "source_record_id": promoted.group(1),
            "source_file": "cases/R2_HUMAN_VALIDATION_uncertainties.csv; cases/FINAL_DATASET_FOR_HUMAN_REVIEW.csv",
            "source_artifact_status": "derived_from_intact_uncertainty_validation",
        }
    return {
        "human_decision": row.get(DECISION_COLUMN, "") or "NOT_RECORDED",
        "human_decision_source": f"cases/R2_HUMAN_VALIDATION_relations.csv:{row['relation_id']} (surviving decision cell)",
        "human_rationale": "NOT_RECORDED",
        "decision_date": "NOT_RECORDED",
        "source_record_id": row["relation_id"],
        "source_file": "cases/R2_HUMAN_VALIDATION_relations.csv; cases/FINAL_DATASET_FOR_HUMAN_REVIEW.csv",
        "source_artifact_status": "reconstructed_from_session_and_surviving_decisions",
    }


def human_status(role_status: str) -> str:
    return "RESEARCHER_APPROVED" if role_status == "CONFIRMED" else role_status


def empty_review_fields() -> dict[str, str]:
    return {
        "final_human_decision": "",
        "final_human_role": "",
        "final_human_mechanism": "",
        "final_human_notes": "",
    }


def render_readme(stats: dict[str, dict[str, object]]) -> str:
    rows = []
    for case_id, stat in stats.items():
        rows.append(
            f"| {case_id} | {stat['relations']} | {stat['actors']} | {stat['R']} | {stat['I']} | {stat['T']} | {stat['intermediaries']} | {stat['unresolved']} |"
        )
    total = stats["TOTAL"]
    return f"""# Final Human Review v2 — Actor-in-Role Architecture

## Purpose

This is the final structured package for researcher adjudication. It is **not** a final published dataset. It transforms the extant relationship evidence into an actor-in-role review surface; it does not perform a new legal-text search or recode the three acts.

## Four analytical levels

1. **Act / regulatory regime** — the documentary and comparative context.
2. **Actor-in-role** — the principal substantive unit: actor × act × role (and, only where needed, context).
3. **Intermediary function / mechanism** — the function and orientation recorded when the role is I.
4. **Regulatory relation / episode** — the coding and evidentiary layer that supports the role classification.

## Main scientific logic

Regulatory relationships establish evidence. Actors-in-role are the primary substantive units for mapping regulatory intermediation. Acts/regimes provide the comparative architecture. R/I/T are contextual roles, not permanent properties of organisations.

## Evidence and recital rule

The relation evidence file preserves operative anchors and quotations available in the source dataset. `NOT_RECOVERED` is retained where the damaged R2 relation CSV no longer preserved a quotation; it is not replaced with a new quotation. Recitals can support interpretation but cannot independently generate a positive relationship.

## Human decisions and provenance

The package distinguishes researcher decisions already recorded in R2, the current agent-derived transformation, and pending final adjudication. The source relation CSV was damaged after an earlier complete read: it is preserved untouched; the 20 base records in `FINAL_DATASET_FOR_HUMAN_REVIEW.csv` were reconstructed from that read plus surviving decision cells; five further relations were promoted from explicit researcher changes in the intact uncertainty file. See `06_PROVENANCE.csv` for the full trail. Missing human rationales are `NOT_RECORDED`, never invented.

## Files

- `01_RELATION_EVIDENCE.csv` — one record per relation/episode, including direct, uncertain, and not-supported records.
- `02_ACTOR_ROLE_MAP.csv` — the master actor × case × role map.
- `03_INTERMEDIARY_MECHANISMS.csv` — one intermediary actor-role × mechanism × case.
- `04_ACT_REGIME_SUMMARY.csv` — one row per act.
- `05_UNRESOLVED_FOR_RESEARCHER.csv` — only issues requiring a substantive researcher decision.
- `06_PROVENANCE.csv` — source, reconstruction, method, prompt, and derivation trail.
- `ACTOR_NORMALIZATION_LOG.csv` — conservative name-normalization decisions.
- `FINAL_HUMAN_REVIEW_WORKBOOK.xlsx` — review workbook; CSVs remain canonical.

## Review instructions

Open `FINAL_HUMAN_REVIEW_WORKBOOK.xlsx`, begin with **UNRESOLVED_CASES**, then review **ACTOR_ROLE_MAP** and **INTERMEDIARY_MECHANISMS**. Fill only the blank `final_human_*` columns. Do not alter evidence or provenance columns. The one pending substantive item is `case_02-U07`; it remains `RESEARCHER_DECISION_REQUIRED` and is not promoted to the accepted data.

## Current derived counts

| Case | Relation evidence | Unique actors | R roles | I roles | T roles | Unique intermediaries | Unresolved issues |
| --- | ---: | ---: | ---: | ---: | ---: | ---: | ---: |
{chr(10).join(rows)}
| **TOTAL** | **{total['relations']}** | **{total['actors']}** | **{total['R']}** | **{total['I']}** | **{total['T']}** | **{total['intermediaries']}** | **{total['unresolved']}** |

These are three-case calibration results, not population estimates.
"""


def write_workbook(tables: dict[str, tuple[list[dict[str, str]], list[str]]]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheet_map = {
        "ACTOR_ROLE_MAP": "02_ACTOR_ROLE_MAP.csv",
        "INTERMEDIARY_MECHANISMS": "03_INTERMEDIARY_MECHANISMS.csv",
        "RELATION_EVIDENCE": "01_RELATION_EVIDENCE.csv",
        "ACT_REGIME_SUMMARY": "04_ACT_REGIME_SUMMARY.csv",
        "UNRESOLVED_CASES": "05_UNRESOLVED_FOR_RESEARCHER.csv",
        "PROVENANCE": "06_PROVENANCE.csv",
    }
    for title, source in sheet_map.items():
        rows, fields = tables[source]
        sheet = workbook.create_sheet(title)
        sheet.append(fields)
        for row in rows:
            sheet.append([row.get(field, "") for field in fields])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for column in range(1, sheet.max_column + 1):
            letter = get_column_letter(column)
            length = max(len(str(sheet.cell(row, column).value or "")) for row in range(1, sheet.max_row + 1))
            sheet.column_dimensions[letter].width = min(max(length + 2, 12), 55)
            for cell in sheet[letter]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[1].height = 32
    workbook.save(OUT / "FINAL_HUMAN_REVIEW_WORKBOOK.xlsx")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    source_rows = read_csv(SOURCE, delimiter=";")
    if len(source_rows) != 25:
        raise RuntimeError(f"Expected 25 consolidated relations, found {len(source_rows)}")
    if len({row['relation_id'] for row in source_rows}) != len(source_rows):
        raise RuntimeError("Duplicate relation_id in source matrix")

    raw_uncertainties = read_csv(UNCERTAINTIES, delimiter=";")
    # The uncertainty CSV has a three-line human-readable preamble. Read it directly
    # at its real header row rather than treating its preamble as data.
    with UNCERTAINTIES.open("r", encoding="utf-8-sig", newline="") as handle:
        records = list(csv.reader(handle, delimiter=";"))
    u_fields = records[3]
    uncertainty_rows = [dict(zip(u_fields, record)) for record in records[4:] if any(record)]
    uncertainty_index = {row["Chave"]: row for row in uncertainty_rows}
    if "case_02-U07" not in uncertainty_index or uncertainty_index["case_02-U07"]["Decisão do pesquisador"]:
        raise RuntimeError("case_02-U07 is no longer the expected unresolved record")

    relation_fields = [
        "relation_id", "case_id", "act_title", "celex", "relation_status", "R_actor", "I_actor", "T_actor",
        "relation_type", "regulatory_action", "regulated_object_or_behavior", "intermediary_function",
        "intermediation_orientation", "mechanism", "operative_anchor_location", "operative_evidence_quote",
        "contextual_support_location", "contextual_support_quote", "third_actor_test_summary", "confidence",
        "researcher_decision", "researcher_rationale", "human_decision", "human_decision_source", "human_rationale",
        "decision_date", "source_record_id", "source_file", "source_artifact_status", "provenance", "human_review_required",
    ]
    relation_rows: list[dict[str, str]] = []
    actor_entries: dict[tuple[str, str, str], dict[str, object]] = {}
    normalization_log: dict[tuple[str, str, str], dict[str, str]] = {}

    for source in source_rows:
        case = source["case_id"]
        decision = decision_metadata(source, uncertainty_index)
        relation_rows.append({
            "relation_id": source["relation_id"], "case_id": case, "act_title": ACTS[case]["title"], "celex": source["celex"],
            "relation_status": relation_status(source), "R_actor": source["R"], "I_actor": source["I"], "T_actor": source["T"],
            "relation_type": source["relation_type"], "regulatory_action": source["action"],
            "regulated_object_or_behavior": source["object"], "intermediary_function": source["mediating_function"],
            "intermediation_orientation": ORIENTATION.get(source["relation_id"], ""), "mechanism": source["mechanism"],
            "operative_anchor_location": source["operative_anchor_location"], "operative_evidence_quote": source["evidence_quote"],
            "contextual_support_location": "NOT_RECOVERED" if source["contextual_support"].startswith("[NOT_RECOVERED") else "",
            "contextual_support_quote": source["contextual_support"], "third_actor_test_summary": source["third_actor_test"],
            "confidence": source["confidence"], "researcher_decision": decision["human_decision"],
            "researcher_rationale": decision["human_rationale"], **decision, "provenance": source["provenance"],
            "human_review_required": "YES" if source["relation_type"] == "uncertain" else "NO",
        })
        if source["relation_type"] == "not_supported":
            continue
        status = "UNCERTAIN" if source["relation_type"] == "uncertain" else "CONFIRMED"
        for role, field in (("R", "R"), ("I", "I"), ("T", "T")):
            if role == "I" and source["relation_type"] != "intermediated":
                continue
            for raw in values(source[field]):
                verbatim = actor_name(raw, role)
                if not verbatim:
                    continue
                normalized, reason = normalize_actor(verbatim)
                log_key = (case, verbatim, normalized)
                normalization_log[log_key] = {
                    "verbatim_name": verbatim, "normalized_name": normalized, "case_id": case,
                    "decision": "NORMALIZED" if verbatim != normalized else "RETAINED", "reason": reason,
                }
                key = (case, normalized, role)
                if key not in actor_entries:
                    actor_entries[key] = {
                        "verbatim": [], "relation_ids": [], "articles": [], "evidence": [], "orientations": [],
                        "functions": [], "mechanisms": [], "status": status, "decisions": [], "decision_sources": [],
                        "rationales": [], "decision_dates": [], "provenance": [],
                    }
                entry = actor_entries[key]
                entry["verbatim"].append(verbatim)
                entry["relation_ids"].append(source["relation_id"])
                entry["articles"].append(source["operative_anchor_location"] or source["source_location"])
                entry["evidence"].append(source["action"])
                entry["decisions"].append(decision["human_decision"])
                entry["decision_sources"].append(decision["human_decision_source"])
                entry["rationales"].append(decision["human_rationale"])
                entry["decision_dates"].append(decision["decision_date"])
                entry["provenance"].append(source["provenance"])
                if status == "CONFIRMED":
                    entry["status"] = "CONFIRMED"
                if role == "I":
                    entry["orientations"].append(ORIENTATION[source["relation_id"]])
                    entry["functions"].append(source["mediating_function"])
                    entry["mechanisms"].append(source["mechanism"])

    def distinct(items: list[str]) -> list[str]:
        return list(dict.fromkeys(item for item in items if item))

    actor_rows: list[dict[str, str]] = []
    role_ids: dict[tuple[str, str, str], str] = {}
    for index, key in enumerate(sorted(actor_entries), start=1):
        case, actor, role = key
        role_ids[key] = f"AR_{case.split('_')[1].upper()}_{index:03d}"
    role_by_actor: defaultdict[tuple[str, str], list[str]] = defaultdict(list)
    for case, actor, role in role_ids:
        role_by_actor[(case, actor)].append(role)
    actor_fields = [
        "actor_role_id", "case_id", "act_title", "celex", "actor_name_verbatim", "actor_name_normalized", "actor_type",
        "governance_level", "role", "role_status", "intermediation_orientation", "primary_intermediary_function",
        "mechanisms_observed", "supporting_relation_ids", "supporting_articles", "operative_evidence_summary", "roles_also_observed",
        "human_decision_basis", "human_decision", "human_decision_source", "human_rationale", "decision_date",
        "human_review_required", "reviewer_notes", "provenance", "final_human_decision",
        "final_human_role", "final_human_mechanism", "final_human_notes",
    ]
    for key in sorted(actor_entries):
        case, actor, role = key
        entry = actor_entries[key]
        type_, level = actor_metadata(actor)
        status = str(entry["status"])
        actor_rows.append({
            "actor_role_id": role_ids[key], "case_id": case, "act_title": ACTS[case]["title"], "celex": ACTS[case]["celex"],
            "actor_name_verbatim": "; ".join(distinct(entry["verbatim"])), "actor_name_normalized": actor,
            "actor_type": type_, "governance_level": level, "role": role, "role_status": status,
            "intermediation_orientation": "; ".join(distinct(entry["orientations"])) if role == "I" else "",
            "primary_intermediary_function": "; ".join(distinct(entry["functions"])) if role == "I" else "",
            "mechanisms_observed": "; ".join(distinct(entry["mechanisms"])) if role == "I" else "",
            "supporting_relation_ids": "; ".join(distinct(entry["relation_ids"])),
            "supporting_articles": "; ".join(distinct(entry["articles"])),
            "operative_evidence_summary": " | ".join(distinct(entry["evidence"])),
            "roles_also_observed": "; ".join(sorted(role_by_actor[(case, actor)])),
            "human_decision_basis": "; ".join(distinct(entry["decisions"])),
            "human_decision": "; ".join(distinct(entry["decisions"])),
            "human_decision_source": " | ".join(distinct(entry["decision_sources"])),
            "human_rationale": " | ".join(distinct(entry["rationales"])),
            "decision_date": "; ".join(distinct(entry["decision_dates"])),
            "human_review_required": "YES" if status != "CONFIRMED" else "NO", "reviewer_notes": "",
            "provenance": "Derived from relation IDs: " + "; ".join(distinct(entry["relation_ids"])), **empty_review_fields(),
        })

    relation_by_id = {row["relation_id"]: row for row in relation_rows}
    mechanism_fields = [
        "intermediary_mechanism_id", "actor_role_id", "case_id", "actor_name_normalized", "mechanism", "function_description",
        "intermediation_orientation", "regulator_actor_or_class", "target_actor_or_class", "supporting_relation_ids", "supporting_articles",
        "operative_evidence_quote_or_summary", "institutionalization_basis", "human_status", "human_review_required", "provenance",
        "human_decision", "human_decision_source", "human_rationale", "decision_date",
        "final_human_decision", "final_human_role", "final_human_mechanism", "final_human_notes",
    ]
    mechanism_groups: dict[tuple[str, str, str], list[dict[str, str]]] = defaultdict(list)
    for row in actor_rows:
        if row["role"] != "I":
            continue
        for mechanism in values(row["mechanisms_observed"]):
            mechanism_groups[(row["actor_role_id"], row["case_id"], mechanism)].append(row)
    mechanism_rows: list[dict[str, str]] = []
    for index, ((actor_role_id, case, mechanism), role_rows) in enumerate(sorted(mechanism_groups.items()), start=1):
        role_row = role_rows[0]
        linked = values(role_row["supporting_relation_ids"])
        relevant = [relation_by_id[relation_id] for relation_id in linked if relation_by_id[relation_id]["mechanism"] == mechanism]
        mechanism_rows.append({
            "intermediary_mechanism_id": f"IM_{case.split('_')[1].upper()}_{index:03d}", "actor_role_id": actor_role_id,
            "case_id": case, "actor_name_normalized": role_row["actor_name_normalized"], "mechanism": mechanism,
            "function_description": " | ".join(distinct([item["intermediary_function"] for item in relevant])),
            "intermediation_orientation": " ; ".join(distinct([item["intermediation_orientation"] for item in relevant])),
            "regulator_actor_or_class": " | ".join(distinct([item["R_actor"] for item in relevant])),
            "target_actor_or_class": " | ".join(distinct([item["T_actor"] for item in relevant])),
            "supporting_relation_ids": "; ".join(item["relation_id"] for item in relevant),
            "supporting_articles": "; ".join(distinct([item["operative_anchor_location"] for item in relevant])),
            "operative_evidence_quote_or_summary": " | ".join(distinct([item["operative_evidence_quote"] if not item["operative_evidence_quote"].startswith("[NOT_RECOVERED") else item["regulatory_action"] for item in relevant])),
            "institutionalization_basis": " | ".join(distinct([item["third_actor_test_summary"] for item in relevant])),
            "human_status": human_status(role_row["role_status"]), "human_review_required": role_row["human_review_required"],
            "provenance": "Derived from " + "; ".join(item["relation_id"] for item in relevant),
            "human_decision": "; ".join(distinct([item["human_decision"] for item in relevant])),
            "human_decision_source": " | ".join(distinct([item["human_decision_source"] for item in relevant])),
            "human_rationale": " | ".join(distinct([item["human_rationale"] for item in relevant])),
            "decision_date": "; ".join(distinct([item["decision_date"] for item in relevant])), **empty_review_fields(),
        })

    summary_fields = [
        "case_id", "act_title", "celex", "legislative_proposer", "formal_adopters", "principal_regulatory_actors", "principal_target_classes",
        "unique_intermediaries_count", "intermediary_actor_list", "mechanisms_observed", "target_facing_intermediaries_count",
        "regulator_facing_intermediaries_count", "bidirectional_intermediaries_count", "direct_relations_count", "intermediated_relations_count",
        "uncertain_relations_count", "governance_architecture_note", "human_review_status",
    ]
    summary_rows: list[dict[str, str]] = []
    stats: dict[str, dict[str, object]] = {}
    for case in ACTS:
        current_relations = [row for row in relation_rows if row["case_id"] == case]
        current_roles = [row for row in actor_rows if row["case_id"] == case]
        intermediaries = [row for row in current_roles if row["role"] == "I"]
        orientations = {row["actor_name_normalized"]: row["intermediation_orientation"] for row in intermediaries}
        mechanisms = Counter(item["mechanism"] for item in mechanism_rows if item["case_id"] == case)
        stats[case] = {
            "relations": len(current_relations), "actors": len({row["actor_name_normalized"] for row in current_roles}),
            "R": sum(row["role"] == "R" for row in current_roles), "I": len(intermediaries),
            "T": sum(row["role"] == "T" for row in current_roles), "intermediaries": len({row["actor_name_normalized"] for row in intermediaries}),
            "unresolved": 1 if case == "case_02" else 0, "mechanisms": dict(mechanisms),
            "target": sum(value == "target-facing" for value in orientations.values()),
            "regulator": sum(value == "regulator-facing" for value in orientations.values()),
            "bidirectional": sum(value == "bidirectional" for value in orientations.values()),
        }
        summary_rows.append({
            "case_id": case, "act_title": ACTS[case]["title"], "celex": ACTS[case]["celex"], "legislative_proposer": "NOT_CODED",
            "formal_adopters": "NOT_CODED", "principal_regulatory_actors": "; ".join(sorted({row["actor_name_normalized"] for row in current_roles if row["role"] == "R"})),
            "principal_target_classes": "; ".join(sorted({row["actor_name_normalized"] for row in current_roles if row["role"] == "T"})),
            "unique_intermediaries_count": str(stats[case]["intermediaries"]),
            "intermediary_actor_list": "; ".join(sorted({row["actor_name_normalized"] for row in intermediaries})),
            "mechanisms_observed": "; ".join(f"{name} ({count})" for name, count in sorted(mechanisms.items())),
            "target_facing_intermediaries_count": str(stats[case]["target"]), "regulator_facing_intermediaries_count": str(stats[case]["regulator"]),
            "bidirectional_intermediaries_count": str(stats[case]["bidirectional"]),
            "direct_relations_count": str(sum(row["relation_type"] == "direct" for row in current_relations)),
            "intermediated_relations_count": str(sum(row["relation_type"] == "intermediated" for row in current_relations)),
            "uncertain_relations_count": str(sum(row["relation_type"] == "uncertain" for row in current_relations)),
            "governance_architecture_note": "Derived only from the retained relation-evidence layer; legislative proposer and formal adopters are not separately coded.",
            "human_review_status": "FINAL_HUMAN_ADJUDICATION_PENDING" if case == "case_02" else "READY_FOR_FINAL_HUMAN_REVIEW",
        })
    stats["TOTAL"] = {key: sum(int(stats[case][key]) for case in ACTS) for key in ("relations", "actors", "R", "I", "T", "intermediaries", "unresolved")}

    unresolved_fields = [
        "issue_id", "case_id", "relation_id", "actor", "proposed_role", "proposed_mechanism", "question", "competing_interpretations",
        "operative_anchor", "evidence_quote", "current_agent_recommendation", "reason_human_decision_required", "final_human_decision", "final_human_notes",
    ]
    u07 = uncertainty_index["case_02-U07"]
    unresolved_rows = [{
        "issue_id": "case_02-U07", "case_id": "case_02", "relation_id": "", "actor": "NOT_IDENTIFIED_IN_CURRENT_CORPUS",
        "proposed_role": "", "proposed_mechanism": "technical assistance / external financing (not coded)",
        "question": "Does the external NDICI-Global Europe technical-assistance mechanism support an actor-in-role classification within this CBAM act?",
        "competing_interpretations": "(1) Keep outside this act: Article 30 only requires reporting on an external instrument. (2) Review the external instrument before deciding whether a separately scoped relationship should be coded.",
        "operative_anchor": "Article 30(1)(f); Article 30(8) (reporting on effects; no operative assistance mechanism in this corpus)",
        "evidence_quote": "the Commission shall evaluate and report on how the financing under that Regulation has contributed to the decarbonisation of the manufacturing industry in LDCs.",
        "current_agent_recommendation": "Do not promote to the accepted relationship or actor-role datasets; retain as RESEARCHER_DECISION_REQUIRED.",
        "reason_human_decision_required": u07["Justificativa preliminar"], "final_human_decision": "", "final_human_notes": "",
    }]

    provenance_fields = ["artifact", "artifact_type", "path", "status", "derived_from", "methodology_version", "prompt_version", "notes"]
    provenance_rows = [
        {"artifact": "R2_HUMAN_VALIDATION_relations.csv", "artifact_type": "human-validation source", "path": "cases/R2_HUMAN_VALIDATION_relations.csv", "status": "PRESERVED_DAMAGED_ORIGINAL", "derived_from": "R2 primary-reader proposals", "methodology_version": "production_v1.1 / v1.2 decisions", "prompt_version": "production_v1.1/1.2", "notes": "CSV delimiter/quoting corruption documented; not repaired in place."},
        {"artifact": "R2_HUMAN_VALIDATION_uncertainties.csv", "artifact_type": "human-validation source", "path": "cases/R2_HUMAN_VALIDATION_uncertainties.csv", "status": "PRESERVED_INTACT", "derived_from": "R2 primary-reader proposals", "methodology_version": "production_v1.1 / v1.2 decisions", "prompt_version": "production_v1.1/1.2", "notes": "Contains 28 decision records, including one unresolved item (case_02-U07)."},
        {"artifact": "FINAL_DATASET_FOR_HUMAN_REVIEW.csv", "artifact_type": "derived relationship matrix", "path": "cases/FINAL_DATASET_FOR_HUMAN_REVIEW.csv", "status": "RECONSTRUCTED_CURRENT_SOURCE", "derived_from": "pre-corruption read + surviving decisions + five explicit uncertainty promotions", "methodology_version": "production_v1.2", "prompt_version": PROMPT_VERSION, "notes": "25 records; source provenance explicitly distinguishes reconstructed base rows from promoted rows."},
        {"artifact": "01_RELATION_EVIDENCE.csv", "artifact_type": "v2 relation-evidence layer", "path": "cases/final_human_review_v2/01_RELATION_EVIDENCE.csv", "status": "DERIVED", "derived_from": "FINAL_DATASET_FOR_HUMAN_REVIEW.csv", "methodology_version": METHOD_VERSION, "prompt_version": PROMPT_VERSION, "notes": "Preserves direct, intermediated, uncertain, and not-supported relationship evidence."},
        {"artifact": "02_ACTOR_ROLE_MAP.csv", "artifact_type": "v2 master actor-role dataset", "path": "cases/final_human_review_v2/02_ACTOR_ROLE_MAP.csv", "status": "DERIVED", "derived_from": "01_RELATION_EVIDENCE.csv", "methodology_version": METHOD_VERSION, "prompt_version": PROMPT_VERSION, "notes": "Principal substantive unit: actor × case × role."},
        {"artifact": "03_INTERMEDIARY_MECHANISMS.csv", "artifact_type": "v2 intermediary-mechanism map", "path": "cases/final_human_review_v2/03_INTERMEDIARY_MECHANISMS.csv", "status": "DERIVED", "derived_from": "02_ACTOR_ROLE_MAP.csv; 01_RELATION_EVIDENCE.csv", "methodology_version": METHOD_VERSION, "prompt_version": PROMPT_VERSION, "notes": "Every row links to a confirmed I actor-role and relation evidence."},
        {"artifact": "frozen legal corpora", "artifact_type": "source corpus", "path": "cases/case_0N/corpus/act.md and act_numbered.md", "status": "PRESERVED_LOCKED", "derived_from": "EUR-Lex source", "methodology_version": "production_v1.1/1.2", "prompt_version": PROMPT_VERSION, "notes": "Consulted only for provenance and unresolved-item verification; not altered."},
    ]

    tables = {
        "01_RELATION_EVIDENCE.csv": (relation_rows, relation_fields),
        "02_ACTOR_ROLE_MAP.csv": (actor_rows, actor_fields),
        "03_INTERMEDIARY_MECHANISMS.csv": (mechanism_rows, mechanism_fields),
        "04_ACT_REGIME_SUMMARY.csv": (summary_rows, summary_fields),
        "05_UNRESOLVED_FOR_RESEARCHER.csv": (unresolved_rows, unresolved_fields),
        "06_PROVENANCE.csv": (provenance_rows, provenance_fields),
        "ACTOR_NORMALIZATION_LOG.csv": (list(normalization_log.values()), ["verbatim_name", "normalized_name", "case_id", "decision", "reason"]),
    }
    for filename, (rows, fields) in tables.items():
        write_csv(OUT / filename, rows, fields)
    (OUT / "README_FINAL_HUMAN_REVIEW.md").write_text(render_readme(stats), encoding="utf-8")
    (OUT / "SUMMARY_STATISTICS.json").write_text(json.dumps(stats, indent=2, ensure_ascii=False), encoding="utf-8")
    write_workbook(tables)

    # Referential and role-integrity checks.
    relation_ids = {row["relation_id"] for row in relation_rows}
    if any(ref not in relation_ids for row in actor_rows for ref in values(row["supporting_relation_ids"])):
        raise RuntimeError("Actor-role record references an unknown relation")
    i_role_ids = {row["actor_role_id"] for row in actor_rows if row["role"] == "I"}
    if any(row["actor_role_id"] not in i_role_ids for row in mechanism_rows):
        raise RuntimeError("Mechanism record references a non-I actor-role")
    if any(not values(row["supporting_relation_ids"]) for row in actor_rows if row["role"] == "I" and row["role_status"] != "RESEARCHER_DECISION_REQUIRED"):
        raise RuntimeError("A supported I actor-role lacks a relationship link")
    for relation in relation_rows:
        r_names = {normalize_actor(actor_name(item, "R"))[0] for item in values(relation["R_actor"])}
        i_names = {normalize_actor(actor_name(item, "I"))[0] for item in values(relation["I_actor"])}
        if r_names & i_names:
            raise RuntimeError(f"R is coded as own I in {relation['relation_id']}")
    print(json.dumps({"output": str(OUT), "stats": stats}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
