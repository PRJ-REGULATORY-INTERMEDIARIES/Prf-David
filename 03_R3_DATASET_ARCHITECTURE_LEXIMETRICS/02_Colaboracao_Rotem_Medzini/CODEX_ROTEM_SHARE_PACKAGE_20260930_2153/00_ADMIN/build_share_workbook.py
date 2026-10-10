"""Build a source-checked discussion workbook from the recovered pilot files.

Only writes inside this package. No historical source file is changed.
"""
from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parents[1]
RUN = ROOT / "AUTONOMOUS_ROTEM_COMPARISON_20260930_2047"
NOW = "2026-09-30"


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def rows(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def write_csv(path: Path, data: list[dict[str, str]], fields: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(data)


def src(sid: str, kind: str, title: str, origin: str, date: str, relpath: str,
        status: str, used: str, url: str = "") -> dict[str, str]:
    p = ROOT / relpath if not re.match(r"^[A-Za-z]:", relpath) else Path(relpath)
    assert p.exists(), p
    return dict(Source_ID=sid, Source_type=kind, Title_or_description=title,
                Author_or_origin=origin, Date=date,
                Path_or_URL=str(p) + (" | " + url if url else ""),
                SHA256=digest(p), Evidence_status=status, Used_for=used)


SOURCES = [
    src("S01", "Correspondence", "Rotem reply, lines 1–20 only", "Rotem Medzini", NOW,
        "Enviados/30_09_26/email e chat", "ROTEM_CONCEPTUAL", "Current methodological positions"),
    src("S02", "Codebook", "Revised Regulatory Intermediation codebook", "Medzini & Levi-Faur project", "2026-09-28",
        "RELATIONAL_LEXIMETRICS_PILOT/PILOT_01/00_SOURCES/Codebook_Regulatory_Intermediation_Revised (1).docx", "ROTEM_DOCUMENTED", "Definitions and documented fields"),
    src("S03", "Codebook", "Unrevised codebook; paragraphs 1–144 only are treated as codebook", "Project baseline", "2026-09",
        "C:/Users/Adm/Downloads/Codebook - Regulatory Intermediation (1).docx", "UNREVISED_BASELINE", "Version context"),
    src("S04", "Chart key", "Chart Key companion to original codebook", "Rotem Medzini", "2026-08-31",
        "C:/Users/Adm/Downloads/Chart Key - Regulatory Intermediation (2).docx", "UNREVISED_BASELINE", "Version context"),
    src("S05", "Dataset", "Regulatory intermediation workbook, six sheets", "Rotem Medzini", "2026-09-02",
        "RELATIONAL_LEXIMETRICS_PILOT/PILOT_01/00_SOURCES/dataset_of_EU_regulatory_intermediation.xlsx", "ROTEM_IMPLEMENTED", "Actual fields and selected case rows"),
    src("S06", "Manuscript", "Law and Governance Architectures, revised v1", "David Levi-Faur", "2026-09-29",
        "RELATIONAL_LEXIMETRICS_PILOT/PILOT_01/00_SOURCES/v1 Law_and_Governance_Architectures_Revised.docx", "CONCEPTUAL_CONTEXT", "Architecture connection only"),
    src("S07", "Human decisions", "PILOT_03 human decision log", "Historical research project", "2026-09",
        "RELATIONAL_LEXIMETRICS_PILOT/PILOT_03/07_HUMAN_ADJUDICATION/HUMAN_DECISION_LOG.csv", "HISTORICAL_HUMAN_ADJUDICATION", "Nineteen decisions and rationales"),
    src("S08", "Impact register", "PILOT_03 backward impact register", "Historical research project", "2026-09",
        "AUTONOMOUS_ROTEM_COMPARISON_20260930_2047/01_BASELINES/BASELINE_REL_PILOT03_BACKWARD_IMPACT_REGISTER.csv", "VERIFIED_HISTORICAL", "Fourteen retrospective flags"),
    src("S09", "Code", "PILOT_03 build script", "Historical research project", "2026-09",
        "RELATIONAL_LEXIMETRICS_PILOT/PILOT_03/99_REPRODUCIBILITY/build_pilot03.py", "VERIFIED_SCRIPT", "A/B provenance"),
    src("S10", "Code", "PILOT_03 incorporation script", "Historical research project", "2026-09",
        "RELATIONAL_LEXIMETRICS_PILOT/PILOT_03/99_REPRODUCIBILITY/incorporate_human_adjudication.py", "VERIFIED_SCRIPT", "Historical gate provenance"),
    src("S11", "Recovered output", "R08 parent representation, first pass only", "Claude comparison run", NOW,
        "AUTONOMOUS_ROTEM_COMPARISON_20260930_2047/08_DUAL_CODING/PARENT_INTERMEDIARY_ACT.csv", "MODEL_FIRST_PASS", "Parent comparison"),
    src("S12", "Recovered output", "R08 provision relation events, first pass only", "Claude comparison run", NOW,
        "AUTONOMOUS_ROTEM_COMPARISON_20260930_2047/08_DUAL_CODING/CHILD_PROVISION_RELATION_EVENT.csv", "MODEL_FIRST_PASS", "Event leads; checked against legal text"),
    src("S13", "Recovered output", "R06 conceptual, documented and implemented alignment matrix", "Claude comparison run", NOW,
        "AUTONOMOUS_ROTEM_COMPARISON_20260930_2047/06_CODEBOOK/ROTEM_CONCEPTUAL_OPERATIONAL_ALIGNMENT_MATRIX.csv", "MODEL_REVIEWED", "Crosswalk leads checked against codebook/workbook"),
    src("S14", "Recovered output", "R07 case register", "Claude comparison run", NOW,
        "AUTONOMOUS_ROTEM_COMPARISON_20260930_2047/07_CASE_SELECTION/CASE_REGISTER.csv", "MODEL_REVIEWED", "Diagnostic case selection"),
    src("S15", "Run record", "Claude run log and confirmed R08 checkpoint", "Claude comparison run", NOW,
        "AUTONOMOUS_ROTEM_COMPARISON_20260930_2047/RUN_LOG.md", "HISTORICAL_RUN_RECORD", "Recovery state"),
]
SOURCES.append(src("S21", "Code", "PILOT_02 build script with constructed base rows", "Historical research project", "2026-09",
    "RELATIONAL_LEXIMETRICS_PILOT/PILOT_02/99_REPRODUCIBILITY/build_pilot02.py", "VERIFIED_SCRIPT", "A/B provenance"))

LEGAL = {
    "CASE_01": ("S16", "32004R0874", "Commission Regulation (EC) No 874/2004", "2004-04-28"),
    "CASE_02": ("S17", "31989L0336", "Council Directive 89/336/EEC", "1989-05-03"),
    "CASE_03": ("S18", "31991L0263", "Council Directive 91/263/EEC", "1991-04-29"),
    "CASE_04": ("S19", "31986L0361", "Council Directive 86/361/EEC", "1986-07-24"),
    "CASE_05": ("S20", "32022R2065", "Regulation (EU) 2022/2065 (Digital Services Act)", "2022-10-19"),
}
for case, (sid, celex, title, date) in LEGAL.items():
    SOURCES.append(src(sid, "Official legal text", title + "; original act English text archived locally",
                       "EU Publications Office / EUR-Lex", date,
                       f"AUTONOMOUS_ROTEM_COMPARISON_20260930_2047/01_BASELINES/official_texts/{celex}.txt",
                       "VERIFIED_OFFICIAL_TEXT", "Provision events and case findings",
                       f"https://eur-lex.europa.eu/legal-content/EN/TXT/?uri=CELEX:{celex}"))
SOURCES.sort(key=lambda r: int(r["Source_ID"][1:]))
SOURCE_IDS = {r["Source_ID"] for r in SOURCES}


def record(fields: list[str], values: tuple | list) -> dict[str, str]:
    assert len(fields) == len(values), (len(fields), len(values))
    return dict(zip(fields, (str(v) if v is not None else "" for v in values)))


CASE_FIELDS = "Case_ID Instrument CELEX_or_reference Provision Intermediary Regulator_or_authority Target_or_recipient Observed_activity Rotem_function Possible_mechanism Recognition_pathway Identification_basis Normative_modality Trigger_condition Relational_issue What_parent_level_preserves What_event_level_adds Conclusion Evidence_status Source_ID".split()
CASES = [
    record(CASE_FIELDS, ("CASE_01", "Regulation (EC) 874/2004", "32004R0874", "Articles 13–14", "Validation agents", "Registry; Commission in selection chain", "Applicants; Registry receives findings", "Receive evidence; examine prior-right claims; notify Registry", "Assurance", "Assurance through documentary verification", "Designation / selection by Registry", "Direct assignment and selection", "Required", "Application and evidence submitted within prescribed deadline", "Applicant → validation agent → Registry has distinct actions and recipients.", "Intermediary identity, function and relevant provisions in the revised schema; descriptive role in the older workbook.", "Links evidence submission, examination and notification to the actor and recipient in Article 14.", "EVENT_LAYER_ADDS_TRACEABILITY", "VERIFIED_OFFICIAL_TEXT; CODEX_CRITERIAL_REVIEW", "S01;S02;S05;S16")),
    record(CASE_FIELDS, ("CASE_02", "Directive 89/336/EEC", "31989L0336", "Annex II, point 5; Article 10(6)", "Insurance provider, inferred from liability insurance requirement", "Member States assess and notify conformity bodies", "Notified body bears insurance requirement; insurer is not assigned an action", "Maintain civil-liability insurance unless State covers liability", "Insurance / risk underwriting in Rotem's workbook", "Risk transfer is a conceptual inference, not a textual insurer action", "None for insurer", "Structural reliance in Rotem's account; qualifying status disputed", "Obligation with exception", "Unless liability is covered by the State under national law", "Does mandatory coverage alone identify an insurer as a regulatory intermediary?", "Names insurers and risk underwriting in implemented coding; revised fields permit rationale and relevant provisions.", "Shows exactly that the obligation is on the body; the insurer is unnamed and has no express action.", "OPEN_CONCEPTUAL_BOUNDARY", "VERIFIED_OFFICIAL_TEXT; CONTESTED_INTERPRETATION", "S01;S02;S05;S07;S17")),
    record(CASE_FIELDS, ("CASE_03", "Directive 91/263/EEC", "31991L0263", "Article 10(1); Annex III, points 3.3–3.4 and 4.3–4.4", "Notified body", "Member States designate; notified body performs assessment", "Manufacturer receives decision and audit or visit reports", "Assess quality system; inspect; audit periodically; make possible unexpected visits/tests; report", "Assurance; monitoring; auditing candidates", "Quality-system assessment; inspection; surveillance audit; reporting", "Vetting in older workbook; legal text uses designation", "Direct assignment with legal designation", "Mixed: shall assess/audit/report; may visit/test", "Periodic audits; unexpected visits and tests subject to conditions", "Several legally distinct activities appear under one intermediary and act.", "Identity, multiple functions and provision list can be retained in revised codebook.", "Connects each activity, output, recipient and modality to a specific Annex III point.", "EVENT_LAYER_ADDS_TRACEABILITY", "VERIFIED_OFFICIAL_TEXT; CODEX_CRITERIAL_REVIEW", "S02;S05;S07;S18")),
    record(CASE_FIELDS, ("CASE_04", "Directive 86/361/EEC", "31986L0361", "Articles 2(8) and 8(3)", "Approved testing laboratory", "Commission", "Commission receives opinion; Member States receive its communication", "Commission consults Member States, then must seek laboratory opinion if no agreement; communicates opinion", "Assurance in Rotem's workbook; expert opinion here is context-specific", "Technical evaluation / opinion; category boundary to discuss", "Designation in older workbook; approval in legal text", "Legally approved role; conditional procedure", "Conditional obligation, then required communication", "Member State decision under 8(1)(a), followed by no agreement within four weeks", "Condition and sequence attach to the Commission's particular request for an opinion.", "Intermediary involvement can be coded Conditional in revised codebook; provision and rationale can document it.", "Links trigger, mandatory request, laboratory opinion and later communication to separate actions.", "EVENT_LAYER_ADDS_TRACEABILITY", "VERIFIED_OFFICIAL_TEXT; HISTORICAL_HUMAN_DECISION_CONSISTENT", "S02;S05;S07;S19")),
    record(CASE_FIELDS, ("CASE_05", "Digital Services Act", "32022R2065", "Article 37(1), (3), (4)", "Independent audit organisation", "Public supervisory framework; no regulator acts within the cited audit event", "Very large online platform/search engine provider", "Conduct independent audit; issue report with opinion and, where needed, recommendations", "Auditing; assurance may also apply under revised definition", "Compliance evaluation and opinion; proposed mechanism label requires calibration", "None in Rotem's implemented row; provider engages organisation", "Direct assignment of audit role", "Annual requirement; recommendations conditional", "Audit opinion not positive triggers operational recommendations", "Auditing is explicit, while assurance is plausible and the mechanism vocabulary overlaps function labels.", "Intermediary identity and audit role are already represented; revised multiple-function field may capture assurance.", "Links audit opinion and conditional recommendations to Article 37(4) without imposing a new function.", "CLARIFICATION_USEFUL", "VERIFIED_OFFICIAL_TEXT; CLASSIFICATION_OPEN", "S01;S02;S05;S20")),
]


EVENT_FIELDS = "Event_ID Case_ID Instrument Provision Text_excerpt Regulator_actor Intermediary_actor Target_or_recipient_actor Action Object Output_or_legal_effect Normative_modality Trigger_condition Function Mechanism_candidate Orientation Evidence_status Source_path_or_URL Coder_provenance Notes".split()
EVENT_INPUTS = [
    ("E01", "CASE_01", "A14", "Every applicant shall submit documentary evidence", "Registry (arrangement)", "Validation agent (request recipient)", "Validation agent", "submit", "documentary evidence of prior right", "Evidence reaches validation agent", "OBLIGATION", "Applicant claims prior right", "N-A: applicant action", "N-A", "N-A", "Action actor is applicant; validation agent is recipient, not actor."),
    ("E02", "CASE_01", "A14", "The relevant validation agent shall examine whether the applicant", "Registry (arrangement)", "Validation agent", "Applicant", "examine", "prior-right claim and evidence", "Finding on prior right", "OBLIGATION", "First in line and timely evidence", "Assurance", "Documentary verification", "target-facing", "The target of examination is applicant's claim."),
    ("E03", "CASE_01", "A14", "he shall notify the Registry of this", "Registry", "Validation agent", "Registry", "notify", "missing/insufficient evidence finding", "Registry informed", "CONDITIONAL_OBLIGATION", "Evidence late or insufficient", "Assurance", "Reporting", "regulator-facing", "Distinct recipient from the examination event."),
    ("E04", "CASE_02", "ANNEX II", "possession of civil liability insurance unless such liability is covered by the State", "Member State (assesses body)", "No insurer action specified", "N-A: the notified body itself bears the condition", "possess", "civil liability insurance", "Eligibility condition for body", "OBLIGATION_WITH_EXCEPTION", "State assumes liability under national law", "N-A: body obligation", "N-A", "N-A", "Action actor is the notified body; text does not name, approve or direct an insurer."),
    ("E05", "CASE_03", "ANNEX III point 3.3", "The notified body shall assess the quality system", "Member State (designation context)", "Notified body", "Manufacturer", "assess", "quality system against requirements", "Assessment decision", "OBLIGATION", "Application for quality-system assessment", "Assurance", "Quality-system assessment", "target-facing", "Historical H_012 supports assurance for point 3.3."),
    ("E06", "CASE_03", "ANNEX III point 3.3", "The evaluation procedure shall include an inspection visit", "Member State (designation context)", "Notified body / auditing team", "Manufacturer", "inspect", "manufacturer's premises", "Evidence for evaluation", "OBLIGATION", "Quality-system evaluation", "Assurance", "Inspection", "target-facing", "Inspection is a component of assessment here."),
    ("E07", "CASE_03", "ANNEX III point 3.4", "The notified body shall evaluate the modifications proposed", "Member State (designation context)", "Notified body", "Manufacturer", "evaluate", "proposed quality-system changes", "Decision whether reassessment required", "OBLIGATION", "Manufacturer proposes update", "Assurance", "Evaluation", "target-facing", "Evaluation names an activity, not necessarily a new function category."),
    ("E08", "CASE_03", "ANNEX III point 4.3", "shall carry out audits at reasonable intervals", "Member State (designation context)", "Notified body", "Manufacturer", "audit", "maintenance and application of quality system", "Audit report to manufacturer", "PERIODIC_OBLIGATION", "Reasonable intervals", "Auditing; monitoring candidate", "Surveillance audit", "target-facing", "Point 4 surveillance supplies context; multiple functions may apply."),
    ("E09", "CASE_03", "ANNEX III point 4.4", "may pay unexpected visits to the manufacturer", "Member State (designation context)", "Notified body", "Manufacturer", "visit", "manufacturer premises", "Visit report", "PERMISSION", "Additional to periodic audit", "Monitoring candidate", "Surveillance visit", "target-facing", "Optional visit, distinct from required periodic audit."),
    ("E10", "CASE_03", "ANNEX III point 4.4", "may carry out, or cause to be carried out, tests", "Member State (designation context)", "Notified body", "Manufacturer", "test", "quality-system functioning", "Test report if performed", "CONDITIONAL_PERMISSION", "During unexpected visit, if necessary", "Assurance or monitoring candidate", "Testing", "target-facing", "Do not infer a test occurred in practice."),
    ("E11", "CASE_04", "A8(3)", "the Commission shall immediately consult the Member States concerned", "Commission", "Approved laboratory not yet involved", "Member States concerned", "consult", "reasons for suspension", "Consultation stage", "CONDITIONAL_OBLIGATION", "Reasons fall under Article 8(1)(a)", "N-A: Commission action", "N-A", "N-A", "First step in sequence."),
    ("E12", "CASE_04", "A8(3)", "the Commission shall seek the opinion of one of the approved laboratories", "Commission", "Approved laboratory (request recipient)", "Approved laboratory", "seek opinion", "technical assessment of dispute", "Laboratory opinion requested", "CONDITIONAL_OBLIGATION", "No agreement within four weeks after consultation", "N-A: Commission action; laboratory function pending", "N-A: technical opinion sought", "N-A", "Commission is the action actor; laboratory is recipient. Historical H_002 confirms the mandatory conditional request."),
    ("E13", "CASE_04", "A8(3)", "The Commission shall communicate the opinion of this laboratory to all the Member States", "Commission", "Approved laboratory (opinion origin)", "Member States", "communicate", "laboratory opinion", "Member States may comment within one month", "OBLIGATION", "After laboratory opinion", "N-A: Commission action", "Information transmission", "N-A", "Do not assign the Commission's communication to the laboratory."),
    ("E14", "CASE_05", "A37(3)", "Audits performed pursuant to paragraph 1 shall be performed by organisations which", "Supervisory framework", "Audit organisation", "Very large platform/search provider", "audit", "compliance with specified obligations/commitments", "Audit assessment", "OBLIGATION", "Audit under Article 37(1), which requires annual provider audits", "Auditing", "Assurance / evaluation candidate", "target-facing", "Article 37(3) identifies the performing organisation; mechanism remains open."),
    ("E15", "CASE_05", "A37(4)", "an audit opinion on whether the provider", "Supervisory framework", "Audit organisation", "Very large platform/search provider", "issue", "audit opinion on compliance", "Positive / positive with comments / negative opinion", "OBLIGATION", "Audit report established", "Auditing; assurance candidate", "Compliance evaluation", "target-facing", "Function and mechanism labels remain interpretive."),
    ("E16", "CASE_05", "A37(4)", "where the audit opinion is not ‘positive’, operational recommendations", "Supervisory framework", "Audit organisation", "Very large platform/search provider", "recommend", "measures and timetable for compliance", "Operational recommendations", "CONDITIONAL_OBLIGATION", "Audit opinion not positive", "Auditing", "Recommendations", "target-facing", "Recommendation is conditional under Article 37(4)(h)."),
]


def pack_section(case: str, provision: str) -> str:
    p = RUN / "07_CASE_SELECTION" / "case_texts" / f"{case}_PROVISIONS.md"
    text = p.read_text(encoding="utf-8")
    celex = LEGAL[case][1]
    candidates = [provision]
    if provision.startswith("ANNEX III point "):
        point = provision.split()[-1]
        candidates = ["ANNEX_III(" + point + ")"]
        if point.startswith("4."):
            candidates.append("ANNEX_III(4)")
    elif provision == "ANNEX II":
        candidates = ["ANNEX_II"]
    for key in candidates:
        match = re.search(r"^## " + re.escape(celex + ":" + key) + r"\s*$", text, re.M)
        if match:
            tail = text[match.end():]
            return re.split(r"^## ", tail, maxsplit=1, flags=re.M)[0]
    raise AssertionError((case, provision))


def norm(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip().casefold()


EVENTS = []
for eid, case, prov, excerpt, regulator, intermediary, target, action, obj, effect, modality, trigger, function, mechanism, orient, note in EVENT_INPUTS:
    section = pack_section(case, prov)
    assert norm(excerpt) in norm(section), (eid, excerpt)
    sid, celex, title, _ = LEGAL[case]
    EVENTS.append(record(EVENT_FIELDS, (eid, case, title, prov, excerpt, regulator, intermediary,
        target, action, obj, effect, modality, trigger, function, mechanism, orient,
        "VERIFIED_OFFICIAL_TEXT; INTERPRETIVE_FIELDS_MODEL_REVIEWED",
        f"https://eur-lex.europa.eu/legal-content/EN/TXT/?uri=CELEX:{celex}",
        "CODEX_SYNTHESIS", note + f" Source_ID: {sid}.")))


CROSS_FIELDS = "Concept_or_variable Rotem_email_position Revised_codebook Workbook_implementation Alignment_status Implication_for_pilot Recommended_action Source_ID".split()
CROSS = [
    ("intermediary × act", "Primary observation: one intermediary per act.", "INTERMEDIARY_MAP has one qualifying category per act; provision is evidence.", "List of Actors/Intermediaries has separate rows; older Coding Digital Intermediaries has act-level packed rows.", "PRESENT_AND_IMPLEMENTED", "Use revised unit as parent; do not treat older packed rows as a revised-schema defect.", "NO_CHANGE_RECOMMENDED"),
    ("regulator_context", "May have several values; denotes institutional context, not actor participation in each activity.", "Multivalued institutional-context field.", "No matching column; older descriptive regulator fields exist.", "DESCRIBED_BUT_NOT_IMPLEMENTED", "Do not form Cartesian actor relationships from values.", "NO_CHANGE_RECOMMENDED"),
    ("target_context", "May have several values; does not identify activity recipients.", "Multivalued target-category field.", "No matching column; target precursors occur in newer sheets.", "DESCRIBED_BUT_NOT_IMPLEMENTED", "Activity recipients require provision-specific reading.", "NO_CHANGE_RECOMMENDED"),
    ("orientation", "No separate position stated in reply.", "Direction toward whom intermediary acts; multiple values possible.", "No orientation column; Position_in_governance_chain is a precursor.", "DESCRIBED_BUT_NOT_IMPLEMENTED", "Optional event links make direction inspectable.", "OPTIONAL_EXTENSION"),
    ("recognition_pathway", "How intermediary is approved or selected; separate from identification basis.", "Vetting; Designation; None; N/A.", "Recognition columns exist, including None for insurer in older coding.", "PRESENT_AND_IMPLEMENTED", "Keep separate from legal basis of identifying an intermediary.", "NO_CHANGE_RECOMMENDED"),
    ("identification basis", "Structural reliance may identify intermediaries; direct assignment/legal recognition boundaries need discussion.", "Inclusion logic is discussed; no separate classification field.", "No dedicated column; explanations can carry basis.", "NOT_FOUND", "Insurer example remains a conceptual boundary.", "DISCUSS_WITH_ROTEM"),
    ("relevant provisions", "Field identifies provisions relevant to the intermediary.", "INTERMEDIARY_MAP lists all relevant provisions.", "List of Actors contains article references; older Coding Digital Intermediaries field is sparsely populated.", "PRESENT_AND_IMPLEMENTED", "Documentary traceability is available in revised schema; activity links may add detail.", "NO_CHANGE_RECOMMENDED"),
    ("rationale", "Several rationale fields explain coding with provision references.", "Value-specific rationale fields appear across codebook.", "No dedicated rationale columns in supplied workbook; narrative columns exist.", "DOCUMENTATION_IMPLEMENTATION_GAP", "Compare documented design and older file by version date.", "NO_CHANGE_RECOMMENDED"),
    ("functions", "At least one regulatory function required under restrictive approach.", "Ten function flags including assurance, monitoring, auditing and insurance.", "Nine corresponding flags in Coding Digital Intermediaries; function text in newer sheets.", "PRESENT_AND_IMPLEMENTED", "Function boundaries require calibration, not presumed schema replacement.", "CLARIFICATION_ONLY"),
    ("mechanism_family", "Functions and mechanisms are distinct.", "Approved mechanism families; family can overlap function names.", "No separate mechanism_family column; Primary_function_mechanism merges wording.", "DESCRIBED_BUT_NOT_IMPLEMENTED", "Keep labels conceptually separate in interpretation.", "CLARIFICATION_ONLY"),
    ("multiple functions", "May apply to one intermediary.", "Multiple positive function flags permitted.", "Multiple Yes flags occur in Coding Digital Intermediaries rows.", "PRESENT_AND_IMPLEMENTED", "Some overlap can be represented already.", "NO_CHANGE_RECOMMENDED"),
    ("assurance", "Named as a calibration topic.", "Assessing/verifying conformity with requirements.", "Function_assurance exists.", "PRESENT_AND_IMPLEMENTED", "Quality-system assessment in Annex III 3.3 is a clear example.", "CLARIFICATION_ONLY"),
    ("monitoring", "Named as a calibration topic.", "Systematic tracking of conduct, compliance, risk, implementation or outcomes.", "Function_monitoring exists.", "PRESENT_AND_IMPLEMENTED", "Surveillance context may coexist with auditing; map activity first.", "CLARIFICATION_ONLY"),
    ("auditing", "Named as a calibration topic.", "Systematic examination identified or structured as an audit.", "Function_auditing exists; older 91/263 row says No despite Annex III 4.3.", "PRESENT_AND_IMPLEMENTED", "The case raises a candidate recoding question for the older row; revised rule itself is clear.", "DISCUSS_WITH_ROTEM"),
    ("evaluation", "Named as a calibration topic.", "No standalone function flag; may describe assurance activity.", "No standalone function column.", "NOT_FOUND", "Do not add a function until an uncaptured governance task is shown.", "CLARIFICATION_ONLY"),
    ("transmission of regulatory information", "Named as a calibration topic.", "No separate function flag; reporting may be described within existing tasks and obligations.", "Reporting language occurs in descriptive fields.", "DOCUMENTATION_IMPLEMENTATION_GAP", "Distinguish transmitting an assessment result from a standalone information function.", "CLARIFICATION_ONLY"),
    ("insurance / risk underwriting", "Insurance requirement can ground structural reliance; recognition is None.", "Insurance / risk-underwriting function flag and mechanism family.", "Insurance companies listed for 89/336 and 91/263 in older Coding Digital Intermediaries sheet.", "PRESENT_AND_IMPLEMENTED", "Historical human H_009/H_010 leaves qualifying status unresolved.", "DISCUSS_WITH_ROTEM"),
    ("intermediary_involvement / modality", "No specific claim beyond existing schema in reply.", "Required; Conditional; Optional; Member State option; Not specified.", "No matching column in supplied workbook.", "DESCRIBED_BUT_NOT_IMPLEMENTED", "Parent can capture conditionality; event link locates exact trigger/action.", "OPTIONAL_EXTENSION"),
]
CROSS = [record(CROSS_FIELDS, x + ("S01;S02;S05;S13",)) for x in CROSS]


FM_FIELDS = "Case_ID Provision Actor Textual_activity Function_candidate Mechanism_candidate Alternative_classification Reason_for_ambiguity Existing_codebook_rule Does_multifunction_solve_it Clarification_needed Proposed_boundary_rule Confidence Source_ID".split()
FM = [
    ("CASE_03", "Annex III 3.3", "Notified body", "Assess quality system; inspect premises", "Assurance", "Quality-system assessment", "Evaluation as activity, not necessarily function", "Assessment and evaluation words occur within assurance task.", "Function_assurance = assess/verify conformity; H_012/H_013 separate point 3.3 from surveillance.", "YES", "Whether inspection is separately coded when embedded in assessment.", "Code assurance for conformity assessment; describe evaluation/inspection as activity.", "HIGH", "S02;S07;S18"),
    ("CASE_03", "Annex III 4.3", "Notified body", "Carry out audits at reasonable intervals", "Auditing; monitoring candidate", "Surveillance audit", "Auditing alone; monitoring alone", "Audit is explicit within a surveillance section.", "Auditing requires audit-structured examination; monitoring is systematic tracking.", "PARTIAL", "When both flags apply to recurring surveillance audits.", "Keep auditing when text says audit; add monitoring only if tracking is separately evidenced.", "MEDIUM", "S02;S05;S18"),
    ("CASE_03", "Annex III 4.4", "Notified body", "Unexpected visit; possible tests", "Monitoring or assurance candidate", "Visit / testing", "Auditing", "Visits may be supervisory; tests verify system functioning if necessary.", "Assurance verifies conformity; monitoring tracks implementation; auditing needs audit structure.", "PARTIAL", "How to classify an unexpected visit with optional testing.", "Do not infer a full audit solely from a visit; retain optionality.", "MEDIUM", "S02;S18"),
    ("CASE_04", "Article 8(3)", "Approved laboratory", "Provide technical opinion after Commission request", "Assurance candidate", "Technical evaluation / opinion", "Information provision", "Opinion evaluates a dispute but is later transmitted by Commission.", "At least one regulatory function needed; assurance covers assessment/verification.", "PARTIAL", "Does this opinion assess conformity under assurance definition?", "Code the laboratory's evaluation separately from Commission's communication.", "MEDIUM", "S01;S02;S07;S19"),
    ("CASE_05", "Article 37(1), (4)", "Audit organisation", "Audit compliance; issue opinion and recommendations", "Auditing; assurance candidate", "Compliance evaluation / audit opinion", "Auditing only", "Audit is explicit; assurance definition also plausibly covers compliance assessment.", "Both function flags are available; mechanism_family uses overlapping labels.", "YES", "Which function is primary and which mechanism label is apt?", "Preserve auditing as explicit function; discuss whether assurance is additional.", "MEDIUM", "S01;S02;S05;S20"),
    ("CASE_01", "Article 14", "Validation agent", "Examine claim then notify Registry", "Assurance", "Documentary verification and reporting", "Transmission as separate function", "Notification communicates an assurance result.", "Assurance exists; no discrete transmission flag.", "PARTIAL", "When reporting is itself a regulatory function.", "Keep assurance for validation; record notification as distinct event.", "MEDIUM", "S01;S02;S16"),
]
FM = [record(FM_FIELDS, x) for x in FM]


TRACE_FIELDS = "Case_ID Coding_decision Relevant_provision_available Rationale_available Actor_reconstructable Action_reconstructable Recipient_reconstructable Function_reconstructable Mechanism_reconstructable Condition_reconstructable Documentary_traceability Inferential_traceability Gap Minimal_fix Source_ID".split()
TRACE = [
    ("CASE_01", "Validation and notification in Article 14", "PARTIAL", "PARTIAL", "YES", "YES", "PARTIAL", "YES", "PARTIAL", "YES", "PARTIAL", "PARTIAL", "Older row describes validation but not which activity has Registry as recipient; revised rationale can cite Article 14.", "Optional provision-linked event rows for submissions, examinations and notification.", "S02;S05;S16"),
    ("CASE_02", "Insurer qualifies through required coverage", "PARTIAL", "PARTIAL", "PARTIAL", "PARTIAL", "NO", "PARTIAL", "PARTIAL", "YES", "PARTIAL", "NO", "Annex II supports body obligation; insurer's actor role and function are inferential and contested.", "Discuss qualification rule with Rotem; retain competing interpretations.", "S01;S02;S05;S07;S17"),
    ("CASE_03", "Notified body assessments versus surveillance audit", "PARTIAL", "PARTIAL", "YES", "PARTIAL", "PARTIAL", "PARTIAL", "PARTIAL", "PARTIAL", "PARTIAL", "PARTIAL", "The older row aggregates functions; specific activity/output links need source reading.", "Use revised per-intermediary rationale and optionally link Annex III events.", "S02;S05;S18"),
    ("CASE_04", "Conditional Commission request for laboratory opinion", "PARTIAL", "PARTIAL", "YES", "PARTIAL", "PARTIAL", "YES", "PARTIAL", "PARTIAL", "PARTIAL", "PARTIAL", "Revised involvement can express Conditional; old file does not link condition to seek-opinion action.", "Cite Article 8(3) in involvement rationale; optional event link for sequence.", "S02;S05;S07;S19"),
    ("CASE_05", "Audit opinion and conditional recommendations", "YES", "PARTIAL", "YES", "YES", "PARTIAL", "YES", "PARTIAL", "PARTIAL", "YES", "PARTIAL", "Article 37 is cited; report content and recipient require provision reading.", "Clarify function/mechanism terms; optional event rows only if comparison needs outputs.", "S02;S05;S20"),
]
TRACE = [record(TRACE_FIELDS, x) for x in TRACE]


ARCH_FIELDS = "Pilot_observation Relevant_dimension_in_David_paper Empirical_example Potential_contribution Caveat Source_ID".split()
ARCH = [
    ("Provision cites a legally structured activity", "Provision as anchor", "Directive 86/361, Article 8(3)", "Anchors a condition and sequence to a legal provision.", "Analytical correspondence, not author endorsement.", "S06;S19"),
    ("Different actors perform different steps", "Actors A and B", "Regulation 874/2004, Article 14", "Makes applicant, validation agent and Registry positions explicit.", "Model-reviewed event interpretation.", "S06;S16"),
    ("One intermediary performs several tasks", "Functions", "Directive 91/263, Annex III", "Shows assessment, audit, visit and report as separate tasks.", "Function boundaries remain open.", "S02;S06;S18"),
    ("Similar work may use different operative devices", "Mechanisms", "DSA Article 37", "Separates audit role from opinion and recommendations.", "Mechanism label is a candidate.", "S02;S06;S20"),
    ("Direction and recipient can vary by activity", "Relationships", "Directive 86/361, Article 8(3)", "Distinguishes Commission→laboratory request from Commission→Member States communication.", "No claim of observed practice.", "S06;S07;S19"),
    ("Events can be nested under a parent unit", "Governance architecture", "All selected cases", "Could operationalize relational connections while keeping intermediary × act comparisons.", "Proposal for discussion, not adopted design.", "S01;S02;S06;S16;S17;S18;S19;S20"),
]
ARCH = [record(ARCH_FIELDS, x) for x in ARCH]


LIMIT_FIELDS = "Limitation_ID Description Affected_result Severity Mitigation Remaining_decision Source_ID".split()
LIMITS = [
    ("L01", "Igor's earlier comments to Rotem were not recovered from disk; only Rotem's reply is used for her position.", "Conceptual crosswalk", "MEDIUM", "Quote/attribute only the reply's own statements.", "Retrieve sent comments if a fuller exchange is needed.", "S01;S15"),
    ("L02", "Supplied workbook was modified on 2 September; revised codebook is dated 28 September and documents fields absent from workbook.", "Implementation comparison", "MEDIUM", "Keep conceptual, documented and implemented layers separate.", "Confirm current working dataset version with Rotem.", "S02;S05;S13"),
    ("L03", "PILOT_03 A/B outputs were built by scripts and are not independent model coders.", "PILOT_03 provenance", "HIGH", "Show only verified historical human decisions; omit agreement percentages as reliability.", "None for this attachment.", "S09;S21"),
    ("L04", "Claude's second pass left no output in the recovered run; any later model review would not be human adjudication.", "R08 comparison", "MEDIUM", "Use criterion-based review of selected legal provisions; mark CODEX_SYNTHESIS.", "Human review may follow discussion.", "S15"),
    ("L05", "Insurance provider identification under structural reliance conflicts with the earlier human relational-edge threshold.", "Directive 89/336, Annex II", "HIGH", "Present Rotem's and historical human interpretations without changing either.", "Discuss qualification threshold with Rotem and human adjudicator.", "S01;S07;S17"),
    ("L06", "Five cases were chosen to diagnose specific coding questions, not sampled statistically.", "All empirical findings", "MEDIUM", "Limit conclusions to observed provisions and candidate refinements.", "Broader testing if generalization is intended.", "S14"),
    ("L07", "A small number of older workbook rows pack multiple intermediaries; they do not implement every revised-codebook field.", "Traceability and case comparisons", "MEDIUM", "Credit the revised per-intermediary design and label older-file limits as version gaps.", "Confirm whether an updated dataset exists.", "S02;S05;S13"),
    ("L08", "PILOT_03 blind-challenge excerpts are paraphrases; one changes Article 8(3) 'shall seek' to 'may seek'.", "Modality example", "HIGH", "Use official legal text for all event rows; retain human H_002 unchanged.", "Correct historical excerpt only through separate authorized governance process.", "S07;S09;S19"),
]
LIMITS = [record(LIMIT_FIELDS, x) for x in LIMITS]


def read_pilot() -> tuple[list[dict[str, str]], int]:
    p = ROOT / "RELATIONAL_LEXIMETRICS_PILOT/PILOT_03/07_HUMAN_ADJUDICATION/HUMAN_DECISION_LOG.csv"
    decisions = rows(p)
    ids = [r["decision_id"] for r in decisions]
    assert len(decisions) == len(set(ids)) == 19
    assert all(r["provenance"] == "HUMAN_ADJUDICATED" and r["status"] == "ADJUDICATED" and r["rationale"] for r in decisions)
    impact = sum(r["backward_impact_flag"] == "YES" for r in decisions)
    assert impact == 14
    impact_rows = rows(ROOT / "AUTONOMOUS_ROTEM_COMPARISON_20260930_2047/01_BASELINES/BASELINE_REL_PILOT03_BACKWARD_IMPACT_REGISTER.csv")
    assert {r["decision_id"] for r in impact_rows} == {r["decision_id"] for r in decisions if r["backward_impact_flag"] == "YES"}
    manifest = json.loads((ROOT / "RELATIONAL_LEXIMETRICS_PILOT/PILOT_03/99_REPRODUCIBILITY/PILOT_03_MANIFEST_FINAL.json").read_text(encoding="utf-8"))
    assert manifest["final_gate"] == "ONTOLOGY_STABLE_WITH_REVISIONS"
    return decisions, impact


PILOT_FIELDS = "Decision_ID Issue_type Prior_state Human_decision Human_rationale Retrospective_impact Provenance Source_file".split()
decisions, impact_count = read_pilot()
PILOT = [record(PILOT_FIELDS, (r["decision_id"], r["issue_type"],
           f"PENDING; Option 1: {r['option_1']} | Option 2: {r['option_2']}",
           f"{r['human_decision']}: {r['corrected_value']}", r["rationale"],
           r["backward_impact_flag"], r["provenance"], "HUMAN_DECISION_LOG.csv (S07)")) for r in decisions]


def setup_data_sheet(wb: Workbook, name: str, fields: list[str], data: list[dict[str, str]],
                     note: str | None = None) -> None:
    ws = wb.create_sheet(name)
    header_row = 3 if note else 1
    if note:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(fields))
        ws.cell(1, 1, note)
        ws.row_dimensions[1].height = 39
        ws.cell(1, 1).font = Font(name="Aptos", size=11, bold=True, color="16425B")
        ws.cell(1, 1).fill = PatternFill("solid", fgColor="E9F2F7")
        ws.cell(1, 1).alignment = Alignment(wrap_text=True, vertical="center")
    for col, key in enumerate(fields, 1):
        cell = ws.cell(header_row, col, key)
        cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E68")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="thin", color="ABC6D4"))
        width = 18 if key.endswith("_ID") else (25 if any(x in key for x in ("Provision", "Status", "Source")) else 37)
        if key in ("Text_excerpt", "Human_rationale", "Description", "Rotem_email_position", "Revised_codebook", "Workbook_implementation", "Notes"):
            width = 49
        if key in ("Event_ID", "Case_ID", "Source_ID", "Decision_ID", "Limitation_ID", "Confidence"):
            width = 16
        ws.column_dimensions[get_column_letter(col)].width = min(width, 49)
    ws.row_dimensions[header_row].height = 32
    for ri, row in enumerate(data, header_row + 1):
        maxlen = 0
        for ci, key in enumerate(fields, 1):
            val = row.get(key, "")
            cell = ws.cell(ri, ci, val)
            cell.font = Font(name="Aptos", size=10, color="243746")
            cell.fill = PatternFill("solid", fgColor="FFFFFF" if ri % 2 == 0 else "F3F7F9")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            maxlen = max(maxlen, len(str(val)))
            if key == "Path_or_URL" and isinstance(val, str) and " | https://" in val:
                # Link official texts to the public act; local archive path and hash stay visible.
                cell.hyperlink = val.split(" | ", 1)[1]
                cell.font = Font(name="Aptos", size=10, color="21618C", underline="single")
        ws.row_dimensions[ri].height = min(84, max(28, 14 * min(6, (maxlen // 70) + 1)))
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(fields))}{max(header_row + 1, len(data) + header_row)}"
    ws.sheet_view.zoomScale = 85
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.print_options.horizontalCentered = True


def overview(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "Overview"
    widths = [11, 46, 34, 47, 24]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.merge_cells("A1:E1")
    ws["A1"] = "Pilot Results — Relational Coding of Regulatory Intermediaries"
    ws["A1"].font = Font(name="Aptos Display", size=19, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="173B55")
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 38
    ws.merge_cells("A2:E2")
    ws["A2"] = "Diagnostic comparison between intermediary × act coding and provision-linked relational extraction"
    ws["A2"].font = Font(name="Aptos", size=12, italic=True, color="345C70")
    ws.row_dimensions[2].height = 26
    meta = [
        ("Run source", "Claude comparison R00–R08 first pass (S15)", "Date", NOW),
        ("Empirical cases", "5 selected diagnostic cases", "PILOT_03 decisions", "19 historical human decisions (S07)"),
        ("Retrospective impact", "14 decisions (S07;S08)", "Workbook status", "Source-checked technical discussion draft"),
    ]
    for ri, (a, b, c, d) in enumerate(meta, 4):
        for ci, value in zip((1, 2, 3, 4), (a, b, c, d)):
            cell = ws.cell(ri, ci, value)
            cell.fill = PatternFill("solid", fgColor="E8F0F4")
            cell.font = Font(name="Aptos", size=10, bold=ci in (1, 3), color="173B55")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[ri].height = 28
    ws.merge_cells("A8:E8")
    ws["A8"] = "Findings from selected provisions"
    ws["A8"].font = Font(name="Aptos", size=12, bold=True, color="173B55")
    for ci, h in enumerate(("ID", "Finding", "Evidence basis", "Implication", "Status"), 1):
        cell = ws.cell(9, ci, h)
        cell.fill = PatternFill("solid", fgColor="1F4E68")
        cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
    findings = [
        ("F1", "Intermediary × act is a useful primary comparative unit.", "Rotem reply and revised codebook; older workbook is mixed by version (S01;S02;S05).", "Retain that parent unit for comparison.", "SUPPORTED_WITH_QUALIFICATION"),
        ("F2", "Provision events preserve actor, action and recipient associations.", "Articles 14, Annex III and 8(3) in selected cases (S16;S18;S19).", "Use a subordinate event layer when the relationship matters.", "SUPPORTED_WITH_QUALIFICATION"),
        ("F3", "Assurance, auditing and monitoring require activity-level calibration in some provisions.", "Directive 91/263 Annex III; DSA Article 37 (S02;S18;S20).", "Clarify boundary examples before changing categories.", "SUPPORTED"),
        ("F4", "Identification basis and recognition pathway answer different questions; insurance remains unsettled.", "Rotem reply; 89/336 Annex II; historical H_009/H_010 (S01;S07;S17).", "Discuss structural-reliance qualification with Rotem.", "OPEN"),
        ("F5", "An activity-linked trigger clarifies conditionality and sequence.", "Directive 86/361 Article 8(3); revised involvement field (S02;S19).", "Keep parent Conditional value; link it to the request event if needed.", "SUPPORTED_WITH_QUALIFICATION"),
    ]
    for ri, line in enumerate(findings, 10):
        for ci, val in enumerate(line, 1):
            cell = ws.cell(ri, ci, val)
            cell.font = Font(name="Aptos", size=10, bold=ci == 1, color="243746")
            cell.fill = PatternFill("solid", fgColor="FFFFFF" if ri % 2 == 0 else "F3F7F9")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[ri].height = 57
    ws.merge_cells("A16:E16")
    ws["A16"] = "Connection to Law and Governance Architectures"
    ws["A16"].font = Font(name="Aptos", size=12, bold=True, color="173B55")
    ws.merge_cells("A17:E19")
    ws["A17"] = ("David Levi-Faur's manuscript treats provisions as textual anchors and governance relationships as connections among actors, functions and mechanisms (S06). "
                 "The selected cases suggest that a provision-linked activity table could empirically represent some of those connections while retaining intermediary × act as the comparative parent. "
                 "For example, Article 8(3) of Directive 86/361 links a conditional Commission request to an approved laboratory and a later communication to Member States (S19). "
                 "This is a proposed analytical connection for discussion, not an endorsement of the pilot by the manuscript's author.")
    ws["A17"].alignment = Alignment(wrap_text=True, vertical="center")
    ws["A17"].font = Font(name="Aptos", size=10, color="243746")
    ws["A17"].fill = PatternFill("solid", fgColor="E8F0F4")
    for ri in (17, 18, 19):
        ws.row_dimensions[ri].height = 23
    ws.merge_cells("A21:E21")
    ws["A21"] = "Status: SUPPORTED = direct selected-case support; SUPPORTED_WITH_QUALIFICATION = support with scope/interpretive limit; OPEN = human discussion needed."
    ws["A21"].font = Font(name="Aptos", size=9, color="456476")
    ws["A21"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[21].height = 27
    ws.freeze_panes = "A9"
    ws.sheet_view.zoomScale = 75
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.print_area = "A1:E21"


def build() -> None:
    # Recover all existing R08 files as immutable copies, including superseded inputs.
    r08 = RUN / "08_DUAL_CODING"
    recovered = OUT / "01_RECOVERED_DATA" / "R08_AS_FOUND"
    if not recovered.exists():
        shutil.copytree(r08, recovered)
    copy_manifest = []
    for p in sorted(r08.rglob("*")):
        if p.is_file():
            q = recovered / p.relative_to(r08)
            assert digest(p) == digest(q)
            copy_manifest.append({"relative_path": str(p.relative_to(r08)), "sha256": digest(p), "bytes": p.stat().st_size})
    write_csv(OUT / "01_RECOVERED_DATA" / "R08_RECOVERY_MANIFEST.csv", copy_manifest,
              ["relative_path", "sha256", "bytes"])

    write_csv(OUT / "SOURCE_REGISTER.csv", SOURCES,
              "Source_ID Source_type Title_or_description Author_or_origin Date Path_or_URL SHA256 Evidence_status Used_for".split())
    wb = Workbook()
    overview(wb)
    setup_data_sheet(wb, "Cases", CASE_FIELDS, CASES)
    setup_data_sheet(wb, "Provision_Events", EVENT_FIELDS, EVENTS)
    setup_data_sheet(wb, "Rotem_Crosswalk", CROSS_FIELDS, CROSS)
    setup_data_sheet(wb, "Function_Mechanism", FM_FIELDS, FM)
    setup_data_sheet(wb, "Traceability", TRACE_FIELDS, TRACE)
    setup_data_sheet(wb, "Pilot03", PILOT_FIELDS, PILOT,
        "PILOT_03 is treated here as an ontology/calibration exercise with human adjudication. Its scripted A/B outputs are not independent inter-model reliability evidence. Historical gate ONTOLOGY_STABLE_WITH_REVISIONS is a project judgment, not statistical validation (S09;S10;S21).")
    setup_data_sheet(wb, "David_Architecture_Link", ARCH_FIELDS, ARCH)
    setup_data_sheet(wb, "Sources", list(SOURCES[0]), SOURCES)
    setup_data_sheet(wb, "Limitations", LIMIT_FIELDS, LIMITS)
    target = OUT / "03_WORKBOOK" / "ROTEM_PILOT_RESULTS_SUMMARY.xlsx"
    wb.save(target)
    print(target)
    print("Cases", len(CASES), "Events", len(EVENTS), "Human decisions", len(PILOT), "Impacts", impact_count)
    print("Recovered R08 files", len(copy_manifest))


if __name__ == "__main__":
    build()
