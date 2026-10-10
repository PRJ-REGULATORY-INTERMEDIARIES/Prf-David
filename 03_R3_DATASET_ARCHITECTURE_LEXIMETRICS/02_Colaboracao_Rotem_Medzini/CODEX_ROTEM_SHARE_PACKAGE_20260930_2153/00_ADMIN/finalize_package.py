"""Write the final, hash-based file index after workbook QA succeeds."""
from __future__ import annotations

import hashlib
from pathlib import Path

from openpyxl import load_workbook

OUT = Path(__file__).resolve().parents[1]
BOOK = OUT / "03_WORKBOOK/ROTEM_PILOT_RESULTS_SUMMARY.xlsx"
QA = OUT / "02_VALIDATION/WORKBOOK_QA.md"
assert BOOK.exists() and QA.exists()
assert QA.read_text(encoding="utf-8").count("PASS —") >= 9
wb = load_workbook(BOOK, read_only=True, data_only=True)
size = BOOK.stat().st_size
sha = hashlib.sha256(BOOK.read_bytes()).hexdigest().upper()
content = f"""# Final file index

## Shareable attachment

- **File:** `{BOOK}`
- **Size:** {size:,} bytes
- **SHA-256:** `{sha}`
- **Sheets:** {len(wb.sheetnames)} — {', '.join(wb.sheetnames)}
- **Empirical cases:** {wb['Cases'].max_row - 1}
- **Verified historical PILOT_03 decision records:** {wb['Pilot03'].max_row - 3}; 14 marked with retrospective impact
- **Validation:** PASS; see `02_VALIDATION/WORKBOOK_QA.md`

## Known limits

The five cases are diagnostic. The supplied workbook predates some revised-codebook fields. PILOT_03 A/B outputs are script-built and provide no independent model reliability estimate. The R08 second pass was unfinished. Insurance-provider qualification remains an open human decision. Earlier Igor comments to Rotem were not found on disk.

No flat CSV was exported: a single table would erase the workbook's case–event–source relationships. This index and the recovery/validation notes are internal supporting material; the Excel workbook is the shareable attachment. No email was sent.
"""
(OUT / "FINAL_FILE_INDEX.md").write_text(content, encoding="utf-8")
print(BOOK, size, sha, len(wb.sheetnames))
