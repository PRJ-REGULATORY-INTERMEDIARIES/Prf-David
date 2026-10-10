from pathlib import Path
import csv
import hashlib
import json
import re
from datetime import datetime

from openpyxl import load_workbook
from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCES = ROOT / "00_SOURCES"
OUT = {
    "design": ROOT / "01_DESIGN",
    "crosswalk": ROOT / "02_CROSSWALK",
    "data": ROOT / "03_PILOT_DATA",
    "diagnostics": ROOT / "04_DIAGNOSTICS",
    "results": ROOT / "05_RESULTS",
}
DATASET = SOURCES / "dataset_of_EU_regulatory_intermediation.xlsx"


def clean(value):
    if value is None:
        return ""
    return str(value).replace("\r", " ").replace("\n", " ").replace("�", "'").strip()


def write_csv(path, headers, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def sha256(path):
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def load_dataset():
    wb = load_workbook(DATASET, read_only=True, data_only=True)
    sheets = {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        sheets[ws.title] = rows
    return sheets


def codebook_definitions():
    defs = {}
    doc = Document(SOURCES / "Codebook_Regulatory_Intermediation_Revised (1).docx")
    for table in doc.tables:
        if len(table.columns) < 2:
            continue
        heads = [clean(c.text) for c in table.rows[0].cells]
        if heads and heads[0] in {"Variable", "Variable family"}:
            for row in table.rows[1:]:
                cells = [clean(c.text) for c in row.cells]
                if cells and cells[0]:
                    defs[cells[0]] = " | ".join(cells[1:])
    return defs


def source_inventory(sheets):
    lines = [
        "# Source inventory",
        "",
        "All source copies in this directory are preserved snapshots. No source file was edited.",
        "",
        "| Copy | Original path | Size (bytes) | Modification time | SHA-256 | Workbook structure |",
        "|---|---|---:|---|---|---|",
    ]
    originals = {
        "v1 Law_and_Governance_Architectures_Revised.docx": r"C:\Users\Adm\Downloads\v1 Law_and_Governance_Architectures_Revised.docx",
        "seidl-kosti-paper-4-with-appendix.pdf": r"C:\Users\Adm\Downloads\seidl-kosti-paper-4-with-appendix.pdf",
        "Codebook_Regulatory_Intermediation_Revised (1).docx": r"C:\Users\Adm\Downloads\Codebook_Regulatory_Intermediation_Revised (1).docx",
        "dataset_of_EU_regulatory_intermediation.xlsx": r"C:\Users\Adm\Downloads\dataset_of_EU_regulatory_intermediation.xlsx",
    }
    for name, original in originals.items():
        copy = SOURCES / name
        st = copy.stat()
        structure = ""
        if copy.suffix.lower() == ".xlsx":
            wb = load_workbook(copy, read_only=True, data_only=True)
            structure = "; ".join(f"{ws.title}={ws.max_row}x{ws.max_column}" for ws in wb.worksheets)
        else:
            structure = "immutable document source"
        lines.append(f"| `{name}` | `{original}` | {st.st_size} | {datetime.fromtimestamp(st.st_mtime).isoformat(timespec='seconds')} | `{sha256(copy)}` | {structure} |")
    lines += [
        "",
        "## Dataset candidate decision",
        "",
        "The codebook itself names `dataset_of_EU_regulatory_intermediation.xlsx` as its dataset file. It was therefore selected as the current Rotem-codebook workbook for this pilot. Other R3/R4 workbooks in the workspace were not substituted for it; they belong to different methodological packages or later validation contexts.",
        "",
        "The workbook has six tabs: `Overview of EU legislation`, `List of intermediaries`, `List of actors`, `Fuller list of EU legislations`, `Coding Digital Intermediaries`, and `List of unused`. The pilot uses the assessed rows in `Coding Digital Intermediaries` and retains the original workbook unchanged.",
    ]
    (SOURCES / "SOURCE_INVENTORY.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def selected_cases(sheets):
    rows = sheets["Coding Digital Intermediaries"]
    headers = [clean(x) for x in rows[0]]
    by_row = {}
    for number, row in enumerate(rows[1:], 1):
        values = list(row) + [None] * (len(headers) - len(row))
        by_row[number] = {headers[i]: values[i] for i in range(len(headers))}
    selected = [
        (1, "P01", "simple assurance arrangement; designated third-party body"),
        (13, "P02", "multiple intermediary categories plus assurance and risk-cover distinction"),
        (18, "P03", "same intermediary family performs assessment, surveillance, audit-like and enforcement-adjacent work"),
        (74, "P04", "clear multi-intermediary architecture with validation, auditing, escrow and dispute resolution"),
        (77, "P05", "structurally difficult assurance/audit/self-regulatory-body boundary"),
    ]
    cases = []
    for row_no, case_id, reason in selected:
        d = by_row[row_no]
        code = clean(d.get("Legislation Code"))
        cases.append({
            "pilot_case_id": case_id,
            "dataset_sheet": "Coding Digital Intermediaries",
            "dataset_row": row_no,
            "legislation_name": clean(d.get("Legislation Name")),
            "legislation_code": code,
            "enactment_year": clean(d.get("Enactment Year")),
            "instrument_type": clean(d.get("Type of Regulatory Instrument")),
            "sector": clean(d.get("Sector")),
            "selection_rationale": reason,
            "selection_status": "MODEL_PILOT",
            "original_dataset_record": f"Coding Digital Intermediaries!row_{row_no}",
            "original_dataset_path": r"C:\Users\Adm\Downloads\dataset_of_EU_regulatory_intermediation.xlsx",
            "source_url": eurlex_url(code),
            "provenance": "HUMAN_EXISTING",
        })
    return cases, by_row


def eurlex_url(code):
    mapping = {
        "76/889/EEC": "31976L0889",
        "89/336/EEC": "31989L0336",
        "91/263/EEC": "31991L0263",
        "EC 874/2004": "32004R0874",
        "1781/2006": "32006R1781",
    }
    return f"https://eur-lex.europa.eu/legal-content/EN/TXT/?uri=CELEX:{mapping.get(code, '')}" if code in mapping else "NOT_RESOLVED"


def build_relationships(cases, by_row):
    rows = []
    evidence = []
    decisions = []

    def add(case, suffix, actor_a, action, obj, function, mechanism, rel_type, actor_b, role_a, role_b, actor_i, provision, normalized, instrument, traceability="NORMALIZED_DATASET_STATEMENT", confidence="MEDIUM", note="", role_i=""):
        rid = f"{case['pilot_case_id']}_REL_{suffix}"
        role_i = role_i or ("intermediary" if actor_i else "")
        rows.append({
            "relationship_id": rid,
            "pilot_case_id": case["pilot_case_id"],
            "act_id": case["legislation_code"],
            "provision_id": f"{case['legislation_code']}::{provision or 'NOT_RECORDED'}",
            "actor_a": actor_a,
            "actor_a_identity": actor_a,
            "actor_a_relational_role": role_a,
            "actor_i": actor_i,
            "actor_i_relational_role": role_i,
            "actor_b": actor_b,
            "actor_b_identity": actor_b,
            "actor_b_relational_role": role_b,
            "legally_expressed_action": action,
            "direction": f"{actor_a} -> {actor_i} -> {actor_b}" if actor_i else f"{actor_a} -> {actor_b}",
            "object": obj,
            "function": function,
            "mechanism": mechanism,
            "instrument_or_legal_device": instrument,
            "relationship_type": rel_type,
            "mandatory_authorized_optional": "NOT_CODED_FROM_SOURCE" if provision == "NOT_RECORDED" else "EXISTING_DATASET_SUMMARY",
            "provenance": "MODEL_PILOT",
            "coding_status": "PROVISIONAL_TRACEABLE_RECONSTRUCTION",
            "confidence": confidence,
            "issue_flag": note,
        })
        evidence.append({
            "evidence_id": f"EV_{rid}",
            "relationship_id": rid,
            "act_id": case["legislation_code"],
            "provision_reference": provision or "NOT_RECORDED_IN_CURRENT_WORKBOOK",
            "source_url_or_identifier": case["source_url"],
            "evidence_excerpt_or_normalized_statement": normalized,
            "evidence_type": traceability,
            "source_layer": "LAYER_A_OBSERVED" if traceability == "SOURCE_TEXT" else "LAYER_A_OBSERVED_AS_RETAINED_DATASET_STATEMENT",
            "existing_dataset_record_id": case["original_dataset_record"],
            "existing_intermediary_record_id": "NOT_SEPARATELY_ASSIGNED_IN_CODING_SHEET",
            "provenance": "HUMAN_EXISTING" if traceability != "SOURCE_TEXT" else "SOURCE_TEXT",
            "confidence": confidence,
            "limitation": "Workbook supplies a normalized statement; direct legal excerpt was not available in the copied workbook." if traceability != "SOURCE_TEXT" else "Direct source text available.",
        })
        decisions.append({
            "decision_id": f"DEC_{rid}",
            "relationship_id": rid,
            "decision_type": "relationship_split",
            "existing_value": "intermediary-level flat coding",
            "pilot_value": rid,
            "provenance": "MODEL_PILOT",
            "basis": "Split a composite coded description only where actor, action, object or recipient changes.",
            "adjudication_status": "NOT_HUMAN_ADJUDICATED",
        })

    c = {x["pilot_case_id"]: x for x in cases}
    d = by_row
    # P01: source workbook has a specific article and a single intermediary category.
    add(c["P01"], "01", "Member State authorities", "designate and notify", "technical apparatus and conformity requirements", "constitution/recognition", "designation", "R_I_T", "manufacturers", "regulator", "target", "Notified bodies", "Article 3(2)", "Bodies notified by each Member State; notified bodies issue conformity marks or certificates replacing the manufacturer's conformity statement.", "notification/designation by Member States", confidence="HIGH")
    add(c["P01"], "02", "Notified bodies", "issue conformity marks or certificates", "technical apparatus and conformity requirements", "assurance", "conformity_attestation", "I_T", "manufacturers", "intermediary", "target", "", "Article 3(2)", "Issue conformity marks or certificates that replace the manufacturer's conformity statement.", "conformity mark/certificate", confidence="HIGH")

    # P02: the current record names three categories and mixes assurance with insurance.
    add(c["P02"], "01", "Member State authorities", "approve, supervise and withdraw approval from", "apparatus conformity", "assurance", "vetting", "R_I_T", "manufacturers", "regulator", "target", "Competent bodies", "NOT_RECORDED", "Competent bodies provide technical reports or certificates; bodies must meet competence, integrity, independence and liability requirements.", "technical report/certificate", note="PROVISION_REFERENCE_MISSING")
    add(c["P02"], "02", "Member State authorities", "approve and withdraw approval from", "apparatus conformity", "assurance", "vetting", "R_I_T", "manufacturers", "regulator", "target", "Notified bodies", "NOT_RECORDED", "Notified bodies issue EC type-examination certificates for specified apparatus and are subject to strong capacity and independence requirements.", "EC type-examination certificate", note="PROVISION_REFERENCE_MISSING")
    add(c["P02"], "03", "Insurance companies", "provide civil-liability coverage", "intermediary organizational and financial governance", "risk underwriting", "insurance", "I_T", "competent/notified bodies", "intermediary", "target", "", "NOT_RECORDED", "Insurance companies provide required civil-liability coverage unless liability is assumed by the State.", "civil-liability insurance", confidence="LOW", note="INTERMEDIARY_STATUS_UNCERTAIN")

    # P03: one existing row contains a dense function bundle for notified bodies and labs.
    add(c["P03"], "01", "Member State authorities", "designate and oversee", "terminal equipment conformity", "assurance", "vetting", "R_I_T", "manufacturers", "regulator", "target", "Notified bodies", "NOT_RECORDED", "Member States designate notified bodies and testing laboratories; notified bodies support manufacturer declarations and conformity procedures.", "designation plus conformity procedure", note="PROVISION_REFERENCE_MISSING")
    add(c["P03"], "02", "Notified bodies", "assess, certify and check products", "terminal equipment", "assurance", "conformity_assessment", "I_T", "manufacturers", "intermediary", "target", "", "NOT_RECORDED", "Notified bodies conduct conformity assessment and product checks; they issue EC type-examination certification.", "EC type-examination and product checks", note="PROVISION_REFERENCE_MISSING")
    add(c["P03"], "03", "Notified bodies", "assess and audit quality systems", "manufacturer quality systems", "monitoring", "audit_surveillance", "I_T", "manufacturers", "intermediary", "target", "", "NOT_RECORDED", "Notified bodies assess quality systems and conduct periodic or unannounced surveillance audits.", "quality-system approval and surveillance", confidence="MEDIUM", note="FUNCTION_MECHANISM_BOUNDARY")
    add(c["P03"], "04", "Designated testing laboratories", "perform conformity tests", "terminal equipment", "assurance", "testing", "I_T", "manufacturers", "intermediary", "target", "", "NOT_RECORDED", "Designated testing laboratories perform conformity testing.", "test laboratory", note="PROVISION_REFERENCE_MISSING")

    # P04: the densest relational architecture in the selected sample.
    add(c["P04"], "01", "European Commission", "oversee and consent to Registry arrangements", ".eu domain registration and administration", "oversight", "public_oversight", "R_I_T", "registrars and registrants", "regulator", "target", "Registry", "Article 12; Article 15", "The Commission sets public-policy rules and constrains the Registry through oversight and consent requirements.", "Commission oversight/consent", confidence="MEDIUM")
    add(c["P04"], "02", "Registry", "accredit and contract with", "domain-name registration services", "constitution/recognition", "accreditation", "R_I_T", "registrants", "regulator", "target", "Accredited registrars", "NOT_RECORDED", "The Registry accredits and contracts with registrars, which implement registration requirements through contractual relations with registrants.", "accreditation and contract", note="PROVISION_REFERENCE_MISSING")
    add(c["P04"], "03", "Registry", "select and instruct", "prior-right claims in domain-name applications", "assurance", "validation", "R_I_T", "domain-name applicants", "regulator", "target", "Validation agents", "Article 13", "Validation agents assess the right claimed for a particular name and examine documentary evidence; the Registry selects them objectively and transparently.", "documentary validation", confidence="HIGH")
    add(c["P04"], "04", "Validation agents", "assess prior-right claims", "prior-right claims in domain-name applications", "assurance", "validation", "I_T", "domain-name applicants", "intermediary", "target", "", "Article 13", "Validation agents verify documentary evidence and examine the applicant's prior rights.", "documentary validation", confidence="HIGH")
    add(c["P04"], "05", "Registry", "appoint and contract with", "phased registration administration", "auditing", "independent_audit", "R_I_T", "Registry administration", "regulator", "target", "Independent auditor", "Article 12", "An independent auditor reviews the fair, appropriate and sound operational and technical administration of the phased registration period and reports findings to the Commission.", "independent audit", confidence="HIGH")
    add(c["P04"], "06", "Independent auditor", "report audit findings to", "phased registration administration", "reporting", "information_transmission", "I_T", "European Commission", "intermediary", "regulator", "", "Article 12", "The independent auditor reports findings to the Commission.", "audit report", confidence="HIGH")
    add(c["P04"], "07", "Registry", "select and contract with", "continuity of domain service", "assurance", "escrow", "R_I_T", "local Internet community", "regulator", "target", "Escrow agents", "Article 15", "The Registry contracts with an escrow agent to support continuity of service in re-delegation or unforeseen circumstances.", "escrow agreement", confidence="HIGH")
    add(c["P04"], "08", "Registry", "select and appoint", "abusive registration disputes and Registry decisions", "dispute resolution", "ADR", "R_I_T", "domain-name holders", "regulator", "target", "ADR providers", "NOT_RECORDED", "ADR providers examine disputes submitted by domain-name holders and the Registry; panelists must be independent and impartial.", "ADR procedure", note="PROVISION_REFERENCE_MISSING")
    add(c["P04"], "09", "ADR providers", "adjudicate", "abusive registration disputes", "dispute resolution", "ADR", "I_T", "domain-name holders", "intermediary", "target", "", "NOT_RECORDED", "ADR providers adjudicate abusive registrations and review Registry decisions.", "ADR panel", note="PROVISION_REFERENCE_MISSING")

    # P05: intentionally conservative because the workbook omits the article reference in the coding row.
    add(c["P05"], "01", "Competent public authorities", "condition an exemption on", "payer information and AML/CFT controls", "assurance", "external_audit", "R_I_T", "non-profit organisations", "regulator", "target", "external auditors or recognized self-regulatory body", "Article 18 (as stated in existing narrative)", "Certain non-profit organisations may be exempt from payer-information requirements where they are subject to reporting obligations and external audit, public supervision or supervision by a recognized self-regulatory body.", "audit/supervision condition", confidence="LOW", note="LEGAL_TEXT_VALIDATION_REQUIRED")
    add(c["P05"], "02", "Non-profit organisations", "report to and submit to audit by", "payer information and related records", "reporting", "information_transmission", "I_T", "external auditors or recognized self-regulatory body", "target", "intermediary", "", "Article 18 (as stated in existing narrative)", "The existing coding narrative identifies auditors or a self-regulatory body to which the non-profit organisation reports and which is connected to audit requirements.", "reporting/audit arrangement", confidence="LOW", note="ALTERNATIVE_INTERMEDIARY_CATEGORY")
    return rows, evidence, decisions


def actors_from_relationships(relationships):
    roles = {}
    for r in relationships:
        for key, role in [("actor_a", r["actor_a_relational_role"]), ("actor_i", r["actor_i_relational_role"]), ("actor_b", r["actor_b_relational_role"])]:
            name = r[key]
            if not name:
                continue
            rec = roles.setdefault(name, {"roles": set(), "cases": set(), "contexts": set()})
            rec["roles"].add(role)
            rec["cases"].add(r["pilot_case_id"])
            rec["contexts"].add(r["act_id"])
    out = []
    for i, (name, rec) in enumerate(sorted(roles.items()), 1):
        out.append({
            "actor_id": f"ACTOR_{i:03d}",
            "actor_identity": name,
            "actor_type": "NOT_NORMALIZED_BEYOND_SOURCE",
            "institutional_context": "; ".join(sorted(rec["contexts"])),
            "relational_roles_observed": "; ".join(sorted(rec["roles"])),
            "pilot_cases": "; ".join(sorted(rec["cases"])),
            "provenance": "MODEL_PILOT",
            "normalization_note": "Role is relationship-specific; identity is retained separately.",
        })
    return out


def crosswalk(sheets, defs):
    fields = []
    for sheet, rows in sheets.items():
        if not rows:
            continue
        headers = [clean(x) for x in rows[0]]
        for field in headers:
            if field and (sheet, field) not in fields:
                fields.append((sheet, field))
    for field in defs:
        if field and ("CODEBOOK", field) not in fields:
            fields.append(("CODEBOOK", field))
    rows = []
    for sheet, field in fields:
        low = field.lower()
        if field in {"case_id", "act_title", "Act_title", "Legislation Name", "Enactment Year", "Sector", "Reference Code", "Downloaded as Pdf", "Assessed?", "Included in the text on governance?", "Included in the text on self-regulation?", "In force, changed, or repealed?", "Information/infrastructure", "Horizontal or vertical", "Type of Regulatory Instrument", "Legislation Code", "Who regulates?", "How is self-regulation operationalised?", "regulatory_object"}:
            treatment, location = "MOVE_TO_ACT_TABLE", "ACTS"
        elif "function_" in low or field in {"Function_description", "Function_summary", "intermediary_object", "intermediary_object_rationale", "What do the intermediaries do?", "How are intermediaries being regulated?", "Intermediary_type (for_profit/industry_association/episdemic)", "Recongnition_pathways (vetting/designation/none)", "Obligation_intensity"}:
            treatment, location = "RETAIN_WITH_REDEFINITION", "RELATIONSHIP_CODING / INTERMEDIARY_REGISTER"
        elif field in {"intermediary_record_id", "Who are the intermediaries?", "Intermediary_legal_label", "Actor_type", "Background_intermediary_family", "Intermediation family", "Intermediary_present"}:
            treatment, location = "MOVE_TO_ACTOR_TABLE", "ACTORS / INTERMEDIARY_REGISTER"
        elif field in {"Relevant provision(s)", "Relevant_articles_addressing_intermediaries", "intermediary_involvement_rationale"}:
            treatment, location = "MOVE_TO_PROVISION_EVIDENCE_TABLE", "PROVISIONS / RELATIONSHIP_EVIDENCE"
        elif field in {"Regulator_type", "EU_level_regulator", "National_level_regulator", "Regulator_in_charge", "Regulatee_name", "Regulatee_type", "Direct_beneficiary_name", "Direct_beneficiary_type", "Position_in_governance_chain", "orientation", "regulator_context", "target_context", "intermediary_involvement"}:
            treatment, location = "MOVE_TO_RELATIONSHIP_TABLE", "RELATIONSHIPS"
        elif "strategy" in low or field in {"public_collective_interests", "Observation_type", "Recognizer_type", "regulatory_architecture"}:
            treatment, location = "DERIVE_FROM_RELATIONSHIPS", "ARCHITECTURE_INFERENCE"
        elif field in {"Primary_function_mechanism", "Recognition_pathway"}:
            treatment, location = "RETAIN_WITH_REDEFINITION", "RELATIONSHIP_CODING"
        elif field in {"Column2", "List of unused"}:
            treatment, location = "UNRESOLVED", "PILOT_ISSUE_REGISTER"
        else:
            treatment, location = "RETAIN_DIRECTLY", "ACTS or ACTORS (context-dependent)"
        definition = defs.get(field, "Dataset header retained; no separate codebook definition found.")
        rationale = {
            "MOVE_TO_ACT_TABLE": "Describes the act or its corpus status and is not inherently relational.",
            "MOVE_TO_ACTOR_TABLE": "Names or characterizes an actor/category; role is not fixed by identity.",
            "MOVE_TO_RELATIONSHIP_TABLE": "The value becomes meaningful only with a directed counterpart and action.",
            "MOVE_TO_PROVISION_EVIDENCE_TABLE": "Legal references must be many-to-many and traceable to relational claims.",
            "RETAIN_WITH_REDEFINITION": "Retain, but attach to a relationship or obligation rather than a flat intermediary row.",
            "DERIVE_FROM_RELATIONSHIPS": "Should be calculated only after observed/coded relationships are established.",
            "UNRESOLVED": "Retain for diagnosis; do not silently delete or redesign in production.",
            "RETAIN_DIRECTLY": "No strong evidence in this pilot that relocation is required.",
        }[treatment]
        rows.append({"source_sheet": sheet, "source_field": field, "current_definition": definition, "pilot_treatment": treatment, "provisional_location": location, "rationale": rationale, "provenance": "MODEL_PILOT"})
    return rows


def write_design(cases, relationships, evidence):
    (OUT["design"] / "RELATIONAL_MODEL_V0_1.md").write_text(f"""# Relational Leximetrics Pilot 01 — Relational Model v0.1

## Scope and status

This is an isolated methodological pilot. It does not modify the production dataset, the official codebook, or the source workbook. It tests whether five already assessed rows in Rotem Medzini's current workbook can be represented as provision-anchored directed relationships.

The candidate unit is a legally constituted relationship, represented provisionally as:

`ACT × PROVISION × OBJECT × ACTOR_A × FUNCTION × MECHANISM × RELATIONSHIP_TYPE × ACTOR_B`

An intermediary is represented as an actor position (`ACTOR_I`) when the source supports a distinct third actor. A provision may support several relationships, and one relationship may require several evidence provisions.

## Analytical layers

**Layer A — OBSERVED.** Act, provision reference, actor identity, legally expressed action, direction, modal character, and supporting text or normalized source statement.

**Layer B — CODED.** Object, function, mechanism, relationship type, institutional type, authority, discretion, accountability, transparency, contestability and related classifications. Every coded value must point to Layer A evidence.

**Layer C — INFERRED.** Regulatory intermediation, R–I–T configurations, strategies, architectures and governance logics. These are derived from combinations of Layer A and B; they are not directly coded from an isolated phrase.

## Entity definitions

| Entity | Provisional definition |
|---|---|
| ACTS | One row per legal act, with stable identifier and act-level attributes. |
| PROVISIONS | One row per addressable article, paragraph, subparagraph, annex or other textual unit. |
| ACTORS | One row per legally identifiable actor/category; identity is distinct from role. |
| RELATIONSHIPS | One row per directed legally constituted connection, with A, optional I and B positions. |
| RELATIONSHIP_EVIDENCE | Many-to-many link between relationship claims and provisions/source statements. |
| RELATIONSHIP_CODING | Coded object, function, mechanism, instrument/device and relational attributes. |
| INTERMEDIARY_REGISTER | Derived/consolidated view of actors that perform intermediary work across relationships. |
| ARCHITECTURE_INFERENCE | Higher-order configurations derived from the relationship layer. |
| CODING_DECISIONS | Provenance for existing, pilot, derived and adjudicated decisions. |

## Preliminary relationship ontology

The pilot distinguishes `R_I_T`, `R_T`, `I_T`, `R_I`, `I_R`, and `ACTOR_TO_ACTOR_UNRESOLVED` as relationship forms. `R_I_T` is used only where a regulator, a distinct intermediary and a target are all supported. A chain is not assumed merely because an actor is institutionally important.

Actions are retained as verb phrases (`designate`, `issue`, `assess`, `audit`, `report`, `adjudicate`, `provide coverage`) before standardization. Functions answer what regulatory work is performed; mechanisms answer how influence is exercised; instruments/devices identify the concrete legal or institutional carrier (certificate, audit, registry, escrow, ADR procedure, etc.).

## Actor identity and relational role

`ACTOR_IDENTITY` is stable within the pilot only as a label for the legally identifiable actor/category. `RELATIONAL_ROLE` is relationship-specific and may be regulator, intermediary, target, beneficiary, recipient or another position. `INSTITUTIONAL_CONTEXT` is stored separately because a public authority, private body, registry or self-regulatory body can occupy different positions in different edges.

## Inference rules

1. Create an observed edge only where the current dataset supplies an action and a legal reference or an explicit limitation is recorded.
2. Split a flat record when actor, action, object, recipient, or legal device changes.
3. Link multiple provisions to one edge rather than duplicating the edge.
4. Do not infer implementation practice, effectiveness, or governance logic from formal legal text.
5. Treat function/mechanism duplication as an issue, not as a reason to silently normalize the taxonomy.
6. Preserve the original value and the pilot value side by side; `MODEL_PILOT` never overwrites `HUMAN_EXISTING`.

## Unresolved conceptual questions

- Whether assurance is a function, a mechanism family, or a result of verification/testing.
- Whether audit is a distinct mechanism when it performs monitoring and assurance simultaneously.
- Whether insurance companies are regulatory intermediaries or supporting institutional devices in the electromagnetic-compatibility case.
- Whether a self-regulatory body named as an alternative to an external auditor should be one actor category or an unresolved branch.
- Whether `R_I_T` should be mandatory for an intermediary claim, or whether `I_T` and `R_I` edges can establish the configuration compositionally.
- Whether the current codebook's act-level self-regulation fields can be derived without importing operational assumptions.

The pilot contains {len(cases)} cases, {len(relationships)} provisional relationship rows, and {len(evidence)} evidence links. These are not production counts and must not be generalized.
""", encoding="utf-8")


def write_diagnostics(cases, relationships, evidence):
    (OUT["diagnostics"] / "FUNCTION_MECHANISM_DIAGNOSTIC.md").write_text("""# Function × Mechanism diagnostic

The distinction is empirically useful but not yet stable. The selected records contain assurance, monitoring, auditing, testing, validation, dispute resolution, information transmission, escrow and risk underwriting. Several current rows combine a broad function flag with a mechanism-like label, while the narrative describes a concrete device that is neither identical to the function nor wholly separate from it.

| Pilot pattern | Evidence in pilot | Diagnostic result |
|---|---|---|
| Assurance / conformity assessment | P01 and P03 | Function can describe the regulatory task; certificate, testing or conformity procedure can be the mechanism/device. |
| Monitoring / audit / surveillance | P03 | Distinction is workable only when the action and object are separated: surveillance tracks continuing compliance; audit examines a quality system. |
| Validation / assurance | P04 | Validation is a concrete procedure/device; assurance is the task performed through documentary validation. |
| Reporting / transmission | P04 and P05 | Reporting is the assigned action/function; information transmission is the mechanism only if movement rather than judgment is the dominant work. |
| Insurance / risk underwriting | P02 | Current flag may describe an institutional support device rather than an intermediary function. Status remains unresolved. |

Explicit duplication flags are retained in `RELATIONSHIP_OBSERVATIONS.csv` as `FUNCTION_MECHANISM_BOUNDARY` or related issue flags. The pilot therefore supports a two-stage test (action → regulatory work → means/device) but does not support adding new production categories.
""", encoding="utf-8")
    (OUT["diagnostics"] / "UNIT_OF_OBSERVATION_DIAGNOSTIC.md").write_text("""# Unit-of-observation diagnostic

The current workbook mixes levels. The codebook states that `ACT_SUMMARY` is act-level, `INTERMEDIARY_MAP` is intermediary-category-level, and provisions are evidence. In the selected coding rows, however, a single row may contain several intermediary categories, several functions and several institutional devices.

The five sampled rows yield {n} provisional relationships. The smallest composite row (P01) yields two edges; the densest (P04) yields nine. This is evidence that an intermediary-level row is not a safe proxy for a single relationship. A normalized relationship layer avoids duplicating act-level fields while permitting re-aggregation to act and intermediary views.

Limitation: the sample is purposive and small. The result is a diagnostic of representational capacity, not a prevalence estimate.
""".format(n=len(relationships)), encoding="utf-8")
    (OUT["diagnostics"] / "ACTOR_ROLE_DIAGNOSTIC.md").write_text("""# Actor identity and role diagnostic

The pilot keeps actor identity and relational role in separate columns. The same named category can be a regulator-facing intermediary in one edge, an intermediary-to-target actor in another, or a target of designation/oversight in a different edge. The .eu case demonstrates this most clearly: the Registry is an institutional context, a regulator of registrars, and the selecting/contracting actor for validation, audit, escrow and ADR providers.

`ACTOR_A` and `ACTOR_B` are sufficient for a basic directed edge, but the optional `ACTOR_I` position is needed to represent R–I–T without collapsing the intermediary into either endpoint. A further position may be needed for recipient/beneficiary or oversight actor when the formal legal chain includes a reporting recipient distinct from the regulator.
""", encoding="utf-8")
    (OUT["diagnostics"] / "TRACEABILITY_DIAGNOSTIC.md").write_text("""# Traceability diagnostic

Traceability improves materially when the relationship row stores a provision reference, an official CELEX URL, an existing dataset record ID, and an evidence link. In this pilot, P01 and P04 have article references in the current workbook narrative; P02, P03 and P05 do not have a structured provision reference in the coding sheet, even though their narratives describe detailed legal work.

Accordingly, evidence is classified as `NORMALIZED_DATASET_STATEMENT` unless a direct source excerpt is present. The pilot does not convert normalized narrative into a verbatim quotation. Rows with missing provision references carry explicit issue flags and are not treated as fully source-verified.

This supports the value of provision anchoring while also showing the cost: source retrieval and article-level validation are prerequisites before scaling.
""", encoding="utf-8")


def write_report(cases, relationships, evidence, crosswalk_rows):
    composite = 4
    report = f"""# PILOT_01_REPORT

## Gate classification

**RELATIONAL_MODEL_SUPPORTED_WITH_REVISIONS**

The pilot supports the relational extension as a diagnostic representation, but only with revisions before scaling: structured provision capture, an explicit evidence-confidence state, a controlled function/mechanism distinction, and a policy for unresolved actor categories. The finding is mixed and does not justify changing the production codebook.

## Scope and sample

The source workbook named by the codebook was selected: `dataset_of_EU_regulatory_intermediation.xlsx`. Five already assessed rows were purposively selected: 76/889/EEC, 89/336/EEC, 91/263/EEC, EC 874/2004 and 1781/2006. The sample covers simple assurance, multiple intermediary categories, ongoing surveillance/audit, a multi-intermediary architecture, and an assurance/audit/self-regulatory boundary case.

The sample is not representative. Four of five selected coding rows contain multiple intermediary categories or alternative categories in one flat record. The pilot reconstructs {len(relationships)} provisional edges from those rows and links each edge to the original workbook record.

## Findings against the required questions

1. **Can current intermediary-level coding be represented without information loss?** Partly. Actor names, functions, devices, provisions, recognition and obligations can be carried forward. Information is not lost if original values and pilot splits are retained. Direct legal text is not available for every selected row, so full legal validation remains incomplete.

2. **One INTERMEDIARY_MAP row: one relationship or several?** Several in the sampled material. A flat row may contain a category bundle, several actions and several devices. The relational representation makes the multiplicity explicit. This is a structural result, not a population estimate.

3. **Does provision anchoring improve traceability?** Yes. P01 and P04 show how article-level references connect an action to a relationship and device. P02/P03/P05 show the cost of missing references: a rich narrative remains difficult to audit.

4. **Are Actor A / Actor B sufficient?** Not for R–I–T. The optional `ACTOR_I` position is necessary when the intermediary is a distinct third actor. Additional recipient/oversight positions may be needed for reporting and consent chains.

5. **Can regulator_context and target_context be separated from identities?** Yes. They are contextual attributes of an edge, not substitutes for actor identities. The same Registry or public authority can occupy different positions across edges.

6. **Does Function × Mechanism work?** It works provisionally but is not yet production-ready. Assurance, validation, testing, monitoring, audit and surveillance can be separated by action/object/device, but current flags sometimes collapse them or use mechanism-like terms as functions.

7. **Which variables are genuinely relational?** Actor A/B/I, orientation, direction, action, object acted upon, function, mechanism, legal device, recipient and relationship type.

8. **Which are act-level?** Act title, instrument type, enactment year, sector, legislation code, corpus status, broad regulatory object and act-level narrative of who regulates.

9. **Which should be derived?** Intermediary status across an act, R–I–T configuration, regulatory strategy, architecture, governance logic and counts of functions/mechanisms.

10. **Can architectures be reconstructed from lower-level relationships?** Plausibly. P04 is the clearest demonstration: Registry, registrars, validation agents, auditor, escrow and ADR edges form a configuration. The inference must remain rule-based and evidence-linked.

11. **Does the proposed sequence work?** Yes as a research workflow: legal text → observed actors/actions/provisions → coded object/function/mechanism/relationship → strategy → architecture → governance logic. The pilot does not test causal or implementation claims.

12. **What is lost by retaining only the intermediary-level unit?** Multiple recipients, separate actions, action order, different objects, reporting direction, legally distinct devices, and the ability to distinguish oversight of the intermediary from work performed by it.

13. **What is duplicated by simply adding relational fields to the flat table?** Act title, sector, regulator context, target context, intermediary name, legal provisions and obligations would repeat across every edge. A normalized evidence and act layer is preferable.

14. **What must change before scaling?** Add structured provision IDs and direct excerpts; define whether function and mechanism are orthogonal; split composite actor cells; formalize `ACTOR_I` plus optional recipient/oversight positions; preserve human/model provenance; and run an intercoder reliability test.

## Evidence and limitations

The pilot is grounded primarily in the existing workbook's normalized statements (`HUMAN_EXISTING`) and the conceptual sources. It does not claim that those statements are verbatim quotations. Missing structured provision references are recorded as issues, not silently repaired. Insurance companies, self-regulatory bodies and enforcement labels are treated cautiously where the source row does not clearly establish a regulatory-intermediary relationship.

No production file was altered. No new mechanism category was introduced into production. All pilot values are labeled `MODEL_PILOT` or `DERIVED_RULE`, and unresolved matters are in `PILOT_01_ISSUE_REGISTER.csv`.

## Recommended Next Step

**A. Refine the schema and conduct PILOT_02**, with a small blind/manual provision-retrieval exercise on the same five cases. PILOT_02 should test whether two independent coders can reproduce the action/function/mechanism splits and whether direct provision excerpts reduce the current uncertainty. Only after that should selected relational elements be considered for integration into the existing codebook.
"""
    (OUT["results"] / "PILOT_01_REPORT.md").write_text(report, encoding="utf-8")


def main():
    for path in OUT.values():
        path.mkdir(parents=True, exist_ok=True)
    sheets = load_dataset()
    source_inventory(sheets)
    cases, by_row = selected_cases(sheets)
    relationships, evidence, decisions = build_relationships(cases, by_row)
    actors = actors_from_relationships(relationships)
    defs = codebook_definitions()
    cw = crosswalk(sheets, defs)

    case_headers = list(cases[0].keys())
    write_csv(OUT["data"] / "PILOT_CASE_REGISTER.csv", case_headers, cases)
    write_csv(OUT["data"] / "RELATIONSHIP_OBSERVATIONS.csv", list(relationships[0].keys()), relationships)
    write_csv(OUT["data"] / "RELATIONSHIP_EVIDENCE.csv", list(evidence[0].keys()), evidence)
    write_csv(OUT["data"] / "ACTOR_REGISTER.csv", list(actors[0].keys()), actors)
    write_csv(OUT["data"] / "RELATIONSHIP_CODING.csv", ["relationship_id", "object", "function", "mechanism", "instrument_or_legal_device", "relationship_type", "provenance", "issue_flag"], [{k: r[k] for k in ["relationship_id", "object", "function", "mechanism", "instrument_or_legal_device", "relationship_type", "provenance", "issue_flag"]} for r in relationships])
    write_csv(OUT["data"] / "CODING_DECISIONS.csv", list(decisions[0].keys()), decisions)

    write_csv(OUT["crosswalk"] / "ROTEM_TO_RELATIONAL_CROSSWALK.csv", list(cw[0].keys()), cw)
    by_treatment = {}
    for row in cw:
        by_treatment[row["pilot_treatment"]] = by_treatment.get(row["pilot_treatment"], 0) + 1
    md = ["# Rotem to relational crosswalk", "", "The CSV is the variable-level register. No source field is deleted; each is assigned a provisional treatment.", "", "| Treatment | Fields |", "|---|---:|"]
    md += [f"| `{k}` | {v} |" for k, v in sorted(by_treatment.items())]
    md += ["", "## Interpretation", "", "Act-level variables move to ACTS; actor and role values separate into ACTORS and RELATIONSHIPS; legal references become many-to-many evidence; strategies and architectures are derived after lower-level coding; unresolved or ambiguous fields remain explicit.", ""]
    (OUT["crosswalk"] / "ROTEM_TO_RELATIONAL_CROSSWALK.md").write_text("\n".join(md), encoding="utf-8")

    write_design(cases, relationships, evidence)
    write_diagnostics(cases, relationships, evidence)

    issues = [
        ["ISSUE_001", "P02/P03", "PROVISION_TRACEABILITY", "Several rich coding rows have no structured relevant provision reference.", "BLOCKS_FULL_LEGAL_VALIDATION", "Retrieve and verify articles/annexes before scaling.", "MODEL_PILOT", "OPEN"],
        ["ISSUE_002", "P03", "FUNCTION_MECHANISM", "Monitoring, audit, surveillance and assurance overlap in the existing record.", "DUPLICATION_RISK", "Keep action, function, mechanism and device separate; do not add a production category yet.", "MODEL_PILOT", "OPEN"],
        ["ISSUE_003", "P02", "ACTOR_STATUS", "Insurance companies may be a supporting device rather than a regulatory intermediary.", "INTERMEDIARY_STATUS_UNCERTAIN", "Retain as an unresolved actor edge and adjudicate against text.", "MODEL_PILOT", "OPEN"],
        ["ISSUE_004", "P05", "ACTOR_STATUS", "Auditor and recognized self-regulatory body appear as alternative categories in one narrative.", "ALTERNATIVE_CATEGORY", "Branch the alternatives or establish an explicit OR relation.", "MODEL_PILOT", "OPEN"],
        ["ISSUE_005", "P04", "ROLE_MULTIPLICITY", "Registry is simultaneously institutional context, regulator of registrars and selecting actor for specialist intermediaries.", "ROLE_NOT_FIXED", "Use relationship-specific roles and retain identity separately.", "MODEL_PILOT", "OPEN"],
        ["ISSUE_006", "P02/P03", "FUNCTION_FLAG", "Current function flags may not match the narrative action set (e.g., enforcement or auditing).", "FLAG_NARRATIVE_MISMATCH", "Do not translate flags directly into relational rows without comparison.", "HUMAN_EXISTING", "OPEN"],
        ["ISSUE_007", "ALL", "EVIDENCE_LAYER", "Current workbook statements are normalized summaries, not consistently verbatim excerpts.", "EVIDENCE_TYPE_MIX", "Add direct excerpt field and evidence type to production candidate schema only after testing.", "MODEL_PILOT", "OPEN"],
        ["ISSUE_008", "ALL", "MODALITY", "Required/authorized/optional status is not consistently structured in selected source rows.", "MODALITY_MISSING", "Capture deontic force at provision level.", "MODEL_PILOT", "OPEN"],
        ["ISSUE_009", "ALL", "ACTOR_POSITION", "A/B alone cannot represent R-I-T and reporting recipients without role collapse.", "POSITIONAL_SCHEMA", "Retain optional ACTOR_I and recipient/oversight positions.", "DERIVED_RULE", "OPEN"],
        ["ISSUE_010", "ALL", "SCALING", "Five purposive cases cannot establish prevalence or intercoder reliability.", "SAMPLE_LIMIT", "Run PILOT_02 with independent duplicate coding.", "MODEL_PILOT", "OPEN"],
    ]
    write_csv(OUT["results"] / "PILOT_01_ISSUE_REGISTER.csv", ["issue_id", "case_scope", "issue_type", "description", "diagnostic", "recommended_handling", "provenance", "status"], [dict(zip(["issue_id", "case_scope", "issue_type", "description", "diagnostic", "recommended_handling", "provenance", "status"], row)) for row in issues])
    write_report(cases, relationships, evidence, cw)

    manifest = {
        "pilot": "RELATIONAL_LEXIMETRICS_PILOT/PILOT_01",
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "source_dataset": str(DATASET),
        "selected_cases": [c["pilot_case_id"] for c in cases],
        "relationship_count": len(relationships),
        "evidence_count": len(evidence),
        "actor_count": len(actors),
        "crosswalk_field_count": len(cw),
        "production_modified": False,
    }
    (ROOT / "99_REPRODUCIBILITY" / "PILOT_MANIFEST.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
