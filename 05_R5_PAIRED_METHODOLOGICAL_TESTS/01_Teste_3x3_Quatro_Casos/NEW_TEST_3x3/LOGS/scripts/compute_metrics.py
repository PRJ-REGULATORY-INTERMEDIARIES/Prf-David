"""NEW TEST 3x3 - deterministic recount of recorded metrics from PRE-EXISTING extraction records.

Read-only on every source. Writes only metrics_computed.json beside this script.
No legal interpretation; counts are taken from existing structured fields.
"""
import json, os, re, collections
import pandas as pd
import openpyxl

ROOT = r"G:\Meu Drive\PRJ-David\Prf-David"
OUT = os.path.join(ROOT, "NEW TEST 3x3", "LOGS", "scripts", "metrics_computed.json")
P = lambda *a: os.path.join(ROOT, *a)
M = {}

# ---------------------------------------------------------------- R3 (frozen 2026-09-21)
R3 = P("R4_CROSS_DOMAIN_VALIDATION", "01_baseline", "R3_HISTORY_STAGING", "PRJ-DAVID-R3", "R3_CLEANROOM")
act = pd.read_csv(os.path.join(R3, "04_actor_register", "ACTOR_REGISTER.csv"))
rel = pd.read_csv(os.path.join(R3, "06_relations", "RELATIONS_ADJUDICATED.csv"))
li = pd.read_csv(os.path.join(R3, "05_law_intermediary", "LAW_INTERMEDIARY_REGISTER.csv"))
r3 = {}
for c, g in rel.groupby("celex_id"):
    sup = g[g.final_classification == "SUPPORTED_INTERMEDIATION"]
    trip = set(zip(sup.R_actor_id, sup.I_actor_id, sup.T_actor_id))
    union_sup = set(sup.R_actor_id.dropna()) | set(sup.I_actor_id.dropna()) | set(sup.T_actor_id.dropna())
    union_all = set(g.R_actor_id.dropna()) | set(g.I_actor_id.dropna()) | set(g.T_actor_id.dropna())
    a = act[act.celex_id == c]
    r3[c] = dict(
        actors_native=int(len(a)),
        actors_specific_entities=int((a.actor_class_or_specific_entity.astype(str).str.upper().str.contains("SPECIFIC")).sum()),
        actors_in_any_relation=len(union_all),
        actors_in_supported_relations=len(union_sup),
        regulators_distinct_all_relations=int(g.R_actor_id.nunique()),
        regulators_distinct_supported=int(sup.R_actor_id.nunique()),
        intermediaries_native=int((li.celex_id == c).sum()),
        targets_distinct_all_relations=int(g.T_actor_id.nunique()),
        targets_distinct_supported=int(sup.T_actor_id.nunique()),
        relations_native=int(len(g)),
        by_class=g.final_classification.value_counts().to_dict(),
        intermediation_native=int(len(sup)),
        rit_triples_supported=len(trip),
    )
for c in act.celex_id.unique():
    r3.setdefault(c, {})
M["R3"] = r3

# ---------------------------------------------------------------- R2 (dataset-eurlex)
R2 = P("Before_R2", "dataset-eurlex", "cases")
CASEMAP = {"case_01": "32021R1119", "case_02": "32023R0956", "case_03": "32023R1115"}


def parse_actors(s):
    """Split an R2 actor field on ';' and strip the trailing role annotation, e.g. 'X (intermediary (verification))' -> 'X'.
    Rule: label = text before the first ' (regulator|intermediary|target' marker; otherwise the whole item."""
    if not isinstance(s, str) or not s.strip():
        return []
    out = []
    for part in s.split(";"):
        part = part.strip()
        m = re.match(r"^(.*?)\s*\((?:regulator|intermediary|target)", part, flags=re.I)
        label = m.group(1).strip() if m else part
        if label:
            out.append(label)
    return out


fin = os.path.join(R2, "final_human_review_v2")
ev = pd.read_csv(os.path.join(fin, "01_RELATION_EVIDENCE.csv"))
am = pd.read_csv(os.path.join(fin, "02_ACTOR_ROLE_MAP.csv"))
r2 = {}
for case, g in ev.groupby("case_id"):
    inter = g[g.relation_type == "intermediated"]
    trip = set()
    union_i = set()
    Ri, Ii, Ti = set(), set(), set()
    for _, row in inter.iterrows():
        R, I, T = parse_actors(row.R_actor), parse_actors(row.I_actor), parse_actors(row.T_actor)
        Ri |= set(R); Ii |= set(I); Ti |= set(T)
        for r_ in R:
            for i_ in I:
                for t_ in T:
                    trip.add((r_, i_, t_))
        union_i |= set(R) | set(I) | set(T)
    allr = set(); allt = set(); alli = set(); unionall = set()
    for _, row in g.iterrows():
        R, I, T = parse_actors(row.R_actor), parse_actors(row.I_actor), parse_actors(row.T_actor)
        allr |= set(R); alli |= set(I); allt |= set(T); unionall |= set(R) | set(I) | set(T)
    a = am[am.case_id == case]
    r2[case] = dict(
        celex=CASEMAP[case],
        relations_native=int(len(g)),
        by_type=g.relation_type.value_counts().to_dict(),
        by_status=g.relation_status.value_counts().to_dict(),
        intermediation_native=int(len(inter)),
        actors_native=int(a.actor_name_normalized.nunique()),
        actor_role_records=int(len(a)),
        role_records={k: int(v) for k, v in a.role.value_counts().items()},
        intermediaries_native=int(a[a.role == "I"].actor_name_normalized.nunique()),
        targets_native=int(a[a.role == "T"].actor_name_normalized.nunique()),
        regulators_native=int(a[a.role == "R"].actor_name_normalized.nunique()),
        rit_triples_intermediated_expanded=len(trip),
        actors_in_intermediated_relations_parsed=len(union_i),
        regulators_in_intermediated_relations=len(Ri),
        intermediaries_in_intermediated_relations=len(Ii),
        targets_in_intermediated_relations=len(Ti),
        distinct_actor_labels_in_relation_fields=len(unionall),
    )
M["R2_final_v2"] = r2

# R2 other versions (counts of relations by type) from JSON proposals
vers = {}
for case in CASEMAP:
    for name in ("CLAUDE_CODING_PROPOSAL", "PRIMARY_INTERPRETIVE_READER_PROPOSAL", "PRIMARY_INTERPRETIVE_READER_PROPOSAL_R2"):
        d = json.load(open(os.path.join(R2, case, "primary_coding", name + ".json"), encoding="utf-8"))
        rels = d["relations"]
        types = collections.Counter(r.get("relation_type") for r in rels)
        Is = set(); Rs = set(); Ts = set()
        for r in rels:
            Rs |= {x["label"] for x in r.get("R", [])}
            Is |= {x["label"] for x in r.get("I", [])}
            Ts |= {x["label"] for x in r.get("T", [])}
        vers[f"{case}|{name}"] = dict(relations=len(rels), by_type=dict(types), intermediated=types.get("intermediated", 0),
                                      distinct_R=len(Rs), distinct_I=len(Is), distinct_T=len(Ts))
M["R2_proposal_versions"] = vers

# R2 sent dataset (Email 01)
wb = openpyxl.load_workbook(P("Enviados", "Email 01", "DATASET_EURLEX_RIT_FINAL_FOR_DAVID.xlsx"), read_only=True, data_only=True)
sent = {}
rows = list(wb["RELATION_EVIDENCE"].iter_rows(values_only=True))
hdr = rows[0]
ix = {h: i for i, h in enumerate(hdr)}
for r in rows[1:]:
    if not r[0]:
        continue
    case = r[ix["case_id"]]
    sent.setdefault(case, collections.Counter())["relations"] += 1
    sent[case][str(r[ix.get("relation_type", 0)])] += 1
rows = list(wb["ACTOR_ROLE_MAP"].iter_rows(values_only=True))
hdr = rows[0]
ix = {h: i for i, h in enumerate(hdr)}
for r in rows[1:]:
    if not r[0]:
        continue
    case = r[ix["case_id"]]
    sent.setdefault(case, collections.Counter())["actor_role_records_" + str(r[ix["role"]])] += 1
M["R2_sent_dataset_email01"] = {k: dict(v) for k, v in sent.items()}

# R2 Stage-4 pilot on Climate Law
ra = pd.read_csv(os.path.join(P("Before_R2", "dataset-eurlex", "03_human", "stage4_attempt_02", "ra_4way"), "RA_4way_comparison.csv"))
M["R2_stage4_climate_law"] = dict(candidates_attempt02=int(len(ra)), RA_result=ra.RA_result.value_counts().to_dict())

# ---------------------------------------------------------------- PoC (Entrega v1-v4)
PO = P("Before_R2", "Estudos Preliminares e testes", "PoC 01")
poc = {}
rel4 = pd.read_csv(os.path.join(PO, "Entrega-04", "data", "rit_relationships_v4.csv"))
sh4 = pd.read_csv(os.path.join(PO, "Entrega-04", "data", "screening_hits_v4.csv"))
for c, g in rel4.groupby("source_act"):
    ids = set(g.R_actor_id) | set(g.I_actor_id) | set(g.T_actor_id)
    poc[c] = dict(v4_relations=len(g), v4_actor_ids=len(ids), v4_I=int(g.I_actor_id.nunique()), v4_R=int(g.R_actor_id.nunique()),
                  v4_T=int(g.T_actor_id.nunique()))
for c, n in sh4.groupby("source_celex").size().items():
    poc.setdefault(c, {})["v4_screening_hits"] = int(n)
v3 = pd.read_csv(os.path.join(PO, "Entrega-02", "data", "rit_relationships_v3.csv"))
for c, g in v3.groupby("source_act"):
    poc.setdefault(c, {})["v3_relations"] = len(g)
v2 = pd.read_csv(os.path.join(P("Before_R2", "Estudos Preliminares e testes", "v1", "PoC", "data"), "rit_coded_sample_v2.csv"))
for c, g in v2.groupby("source_act"):
    poc.setdefault(c, {})["v2_sample_rows"] = len(g)
import csv
with open(os.path.join(P("Before_R2", "Estudos Preliminares e testes", "v1", "PoC", "data"), "rit_coded_sample.csv"), encoding="utf-8-sig", newline="") as fh:
    v1rows = list(csv.reader(fh))
    poc["_v1_rows"] = max(0, len(v1rows) - 1)  # header excluded; file has ragged rows so pandas cannot parse it
    poc["_v1_by_act"] = dict(collections.Counter(r[0] for r in v1rows[1:] if r))
M["PoC"] = poc

# ---------------------------------------------------------------- PILOT_01-03
PL = P("RELATIONAL_LEXIMETRICS_PILOT")
A = pd.read_csv(os.path.join(PL, "PILOT_03", "04_MODEL_A", "RELATIONSHIPS_MODEL_A.csv"))
B = pd.read_csv(os.path.join(PL, "PILOT_03", "05_MODEL_B", "RELATIONSHIPS_MODEL_B.csv"))
PR = pd.read_csv(os.path.join(PL, "PILOT_03", "03_SOURCE_PROVISIONS", "PROVISION_REGISTER.csv"))
norm = lambda s: re.sub(r"\s+", " ", str(s).strip().lower())
p3 = {}
for c, g in A.groupby("celex"):
    gb = B[B.celex == c]
    acts = set()
    for col in ("actor_a", "actor_b", "recipient_actor", "oversight_actor"):
        acts |= {norm(x) for x in g[col].dropna()}
    ints = {norm(x) for x in g[g.intermediary_status_derived.isin(["YES", "DERIVED_POSITION"])].actor_b.dropna()}
    regs = {norm(x) for x in g[g.actor_a_role.astype(str).str.lower() == "regulator"].actor_a.dropna()} | {norm(x) for x in g[g.actor_b_role.astype(str).str.lower() == "regulator"].actor_b.dropna()}
    tgts = {norm(x) for x in g[g.actor_a_role.astype(str).str.lower() == "target"].actor_a.dropna()} | {norm(x) for x in g[g.actor_b_role.astype(str).str.lower() == "target"].actor_b.dropna()}
    p3[c] = dict(regulators_A=len(regs), targets_A=len(tgts), provisions=int((PR.celex == c).sum()), edges_modelA=int(len(g)), edges_modelB=int(len(gb)),
                 derived_rit_candidates_A=int((g.derived_rit_candidate == "YES").sum()),
                 derived_rit_candidates_B=int((gb.derived_rit_candidate == "YES").sum()),
                 distinct_actors_A=len(acts), distinct_intermediary_actor_b_A=len(ints),
                 intermediary_status_direct_YES_A=int((g.intermediary_status_direct == "YES").sum()))
M["PILOT_03"] = p3
o1 = pd.read_csv(os.path.join(PL, "PILOT_01", "03_PILOT_DATA", "RELATIONSHIP_OBSERVATIONS.csv"))
M["PILOT_01"] = {k: dict(relationships=int(len(g)), with_actor_i=int(g.actor_i.notna().sum())) for k, g in o1.groupby("act_id")}
c2 = pd.read_csv(os.path.join(PL, "PILOT_02", "02_CODING_A", "RELATIONSHIPS_CODER_A.csv"))
M["PILOT_02"] = {k: dict(relationships=int(len(g)), with_actor_i=int(g.actor_i.notna().sum())) for k, g in c2.groupby("celex")}

# ---------------------------------------------------------------- R4 structural
R4 = P("ENVIO_R4_TO_DAVID")
ex = pd.read_csv(os.path.join(R4, "03_SOURCE_DATA_AND_CORPUS", "R4_STRUCTURAL_EXTRACTION.csv"))
ac = pd.read_csv(os.path.join(R4, "04_R4_2_STRUCTURAL_EXTRACTION", "R4_DIGITAL_ACTOR_CANDIDATE_REGISTER.csv"))
la = pd.read_csv(os.path.join(R4, "04_R4_2_STRUCTURAL_EXTRACTION", "R4_LEGAL_ACTION_REGISTER.csv"))
r4 = {}
for c, g in ex.groupby("CASE_ID"):
    r4[c] = dict(structural_records=int(len(g)), by_source_type=g.SOURCE_TYPE.value_counts().to_dict(),
                 actor_candidates=int((ac.CASE_ID == c).sum()), legal_action_rows=int((la.CASE_ID == c).sum()),
                 distinct_actor_normalized_in_records=int(g.ACTOR_NORMALIZED_PROVISIONAL.nunique()))
M["R4"] = r4

# ---------------------------------------------------------------- GDPR_PILOT_01 (extraction CSV only)
G = P("RGAM", "pilot", "GDPR_PILOT_01", "03_EXTRACTION")
q = pd.read_csv(os.path.join(G, "QUALIFIED_RELATIONS.csv"))
n = pd.read_csv(os.path.join(G, "NORMATIVE_ASSIGNMENTS.csv"))
pdv = pd.read_csv(os.path.join(G, "POSITIONS_AND_DEVICES.csv"))
M["GDPR_PILOT_01"] = dict(qualified_relations=int(len(q)), normative_assignments=int(len(n)), positions_and_devices=int(len(pdv)),
                          distinct_origin=int(q.origin.nunique()), distinct_counterpart=int(q.counterpart.nunique()),
                          distinct_assignment_subjects=int(n.subject.nunique()),
                          distinct_origin_or_counterpart_union=int(len(set(q.origin.dropna()) | set(q.counterpart.dropna()))))

# ---------------------------------------------------------------- Rotem external dataset
wb = openpyxl.load_workbook(P("RELATIONAL_LEXIMETRICS_PILOT", "PILOT_01", "00_SOURCES", "dataset_of_EU_regulatory_intermediation.xlsx"),
                            read_only=True, data_only=True)
cdi = list(wb["Coding Digital Intermediaries"].iter_rows(values_only=True))
hdr = cdi[0]
rot = []
for r in cdi[1:]:
    if r[1] or r[12]:
        rot.append(dict(title=str(r[1]).strip() if r[1] else "", year=str(r[6]).strip() if r[6] else "", code=str(r[12]).strip() if r[12] else "",
                        sector=str(r[10]).strip() if r[10] else "", assessed=str(r[3]).strip() if r[3] else "",
                        intermediaries=str(r[15]).strip()[:200] if r[15] else ""))
M["Rotem_CDI_rows"] = rot
loa = list(wb["List of actors"].iter_rows(values_only=True))
lo = collections.Counter(); last = None
for r in loa[1:]:
    if r[4]:
        last = str(r[4]).strip()
    if r[15]:
        lo[last] += 1
loi = list(wb["List of intermediaries"].iter_rows(values_only=True))
li_ = collections.Counter(); last = None
for r in loi[1:]:
    if r[4]:
        last = str(r[4]).strip()
    if r[15]:
        li_[last] += 1
M["Rotem_LoA_rows"] = dict(lo)
M["Rotem_LoI_rows"] = dict(li_)

wbx = openpyxl.load_workbook(P("CODEX_ROTEM_SHARE_PACKAGE_20260930_2153", "03_WORKBOOK", "ROTEM_PILOT_RESULTS_SUMMARY.xlsx"), read_only=True, data_only=True)
ev_rows = list(wbx["Provision_Events"].iter_rows(values_only=True))
hdr = ev_rows[0]; ix = {h: i for i, h in enumerate(hdr)}
cnt = collections.Counter(r[ix["Case_ID"]] for r in ev_rows[1:] if r and len(r) > ix["Case_ID"] and r[ix["Case_ID"]])
M["Codex_package_events_by_case"] = dict(cnt)
cs = list(wbx["Cases"].iter_rows(values_only=True))
M["Codex_package_cases"] = [dict(case=r[0], instrument=r[1], celex=r[2]) for r in cs[1:] if r and r[0] and str(r[0]).startswith("CASE_")]
with open(P("AUTONOMOUS_ROTEM_COMPARISON_20260930_2047", "01_BASELINES", "BASELINE_R4_EXTRACTION_DSA_ART21_22_37_42.csv"), encoding="utf-8-sig", newline="") as fh:
    M["R4_DSA_art21_22_37_42_rows"] = max(0, len(list(csv.reader(fh))) - 1)

json.dump(M, open(OUT, "w", encoding="utf-8"), indent=1, ensure_ascii=False, default=str)
print("written", OUT)
